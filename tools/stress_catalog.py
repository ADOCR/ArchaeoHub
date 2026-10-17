"""Prueba local de esfuerzo con datos reales; los directorios originales son entradas de lectura.

La entrega se copia físicamente, sin enlaces duros, antes de invocar su dominio.
Ninguna solicitud descarta originales. Los informes y cambios quedan en workdir.
"""
from pathlib import Path
import argparse
import collections
import ctypes
from ctypes import wintypes
import hashlib
import json
import math
import os
import queue
import statistics
import subprocess
import threading
import time
from verify_release import clean_environment


def digest(path, destination=None):
    h=hashlib.sha256()
    if destination is not None:
        destination.parent.mkdir(parents=True,exist_ok=True)
    output=destination.open("xb") if destination is not None else None
    try:
        with path.open("rb") as stream:
            while chunk:=stream.read(4*1024*1024):
                h.update(chunk)
                if output is not None:output.write(chunk)
    finally:
        if output is not None:output.close()
    return h.hexdigest()


class ProcessMemory:
    def __init__(self,pid):
        class Memory(ctypes.Structure):
            _fields_=[("cb",wintypes.DWORD),("PageFaultCount",wintypes.DWORD)]+[(name,ctypes.c_size_t) for name in ("PeakWorkingSetSize","WorkingSetSize","QuotaPeakPagedPoolUsage","QuotaPagedPoolUsage","QuotaPeakNonPagedPoolUsage","QuotaNonPagedPoolUsage","PagefileUsage","PeakPagefileUsage","PrivateUsage")]
        self.struct=Memory;self.kernel=ctypes.WinDLL("kernel32",use_last_error=True);self.psapi=ctypes.WinDLL("psapi",use_last_error=True)
        self.kernel.OpenProcess.argtypes=[wintypes.DWORD,wintypes.BOOL,wintypes.DWORD];self.kernel.OpenProcess.restype=wintypes.HANDLE
        self.kernel.CloseHandle.argtypes=[wintypes.HANDLE]
        self.psapi.GetProcessMemoryInfo.argtypes=[wintypes.HANDLE,ctypes.POINTER(Memory),wintypes.DWORD]
        self.handle=self.kernel.OpenProcess(0x0400|0x0010,False,pid);self.peak=0;self.private_peak=0;self.samples=0
        self.stop=threading.Event();self.thread=threading.Thread(target=self.run,daemon=True);self.thread.start()
    def run(self):
        while not self.stop.is_set():
            item=self.struct();item.cb=ctypes.sizeof(item)
            if self.handle and self.psapi.GetProcessMemoryInfo(self.handle,ctypes.byref(item),item.cb):
                self.peak=max(self.peak,item.PeakWorkingSetSize);self.private_peak=max(self.private_peak,item.PrivateUsage);self.samples+=1
            self.stop.wait(0.1)
    def close(self):
        self.stop.set();self.thread.join(timeout=1)
        if self.handle:self.kernel.CloseHandle(self.handle)
        return {"peakWorkingSetBytes":self.peak,"peakPrivateBytes":self.private_peak,"samples":self.samples}


class Worker:
    def __init__(self,exe):
        self.process=subprocess.Popen([str(exe)],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,env=clean_environment(),creationflags=subprocess.CREATE_NO_WINDOW)
        self.memory=ProcessMemory(self.process.pid);self.messages=queue.Queue();self.stderr=[];self.counter=0;self.timings=collections.defaultdict(list)
        def read():
            for line in self.process.stdout:self.messages.put(line)
            self.messages.put(None)
        def errors():
            for line in self.process.stderr:
                if len(self.stderr)<100:self.stderr.append(line.decode("utf-8",errors="replace").strip())
        threading.Thread(target=read,daemon=True).start();threading.Thread(target=errors,daemon=True).start()
    def rpc(self,op,expected=True,**values):
        self.counter+=1;request={"schemaVersion":1,"requestId":self.counter,"op":op,**values};start=time.perf_counter()
        self.process.stdin.write((json.dumps(request,ensure_ascii=False)+"\n").encode("utf-8"));self.process.stdin.flush()
        raw=self.messages.get(timeout=90)
        if raw is None:raise RuntimeError(f"Worker closed during {op}: {self.stderr}")
        reply=json.loads(raw);elapsed=time.perf_counter()-start;self.timings[op].append(elapsed)
        assert reply["requestId"]==self.counter,reply
        assert reply["ok"]==expected,reply
        return reply.get("result",reply),elapsed
    def close(self):
        memory=self.memory.close();self.process.stdin.close()
        try:code=self.process.wait(timeout=15)
        except subprocess.TimeoutExpired:self.process.kill();self.process.wait(timeout=5);raise
        assert code==0,(code,self.stderr)
        return memory


def distribution(values):
    ordered=sorted(values)
    return {"count":len(values),"medianMs":round(statistics.median(ordered)*1000,2),"p95Ms":round(ordered[max(0,math.ceil(len(ordered)*0.95)-1)]*1000,2),"maxMs":round(max(ordered)*1000,2)} if values else {}


def gui_probe(executable,root,source,output,site,operation):
    user=ctypes.WinDLL("user32",use_last_error=True)
    callback_type=ctypes.WINFUNCTYPE(wintypes.BOOL,wintypes.HWND,wintypes.LPARAM)
    user.EnumWindows.argtypes=[callback_type,wintypes.LPARAM]
    user.GetWindowThreadProcessId.argtypes=[wintypes.HWND,ctypes.POINTER(wintypes.DWORD)]
    user.IsWindowVisible.argtypes=[wintypes.HWND]
    user.SendMessageTimeoutW.argtypes=[wintypes.HWND,wintypes.UINT,wintypes.WPARAM,wintypes.LPARAM,wintypes.UINT,wintypes.UINT,ctypes.POINTER(ctypes.c_size_t)]
    user.SendMessageTimeoutW.restype=ctypes.c_size_t
    profile=root/("gui-profile-"+str(time.time_ns()));report=profile.with_suffix(".json")
    process=subprocess.Popen([str(executable),"--data-dir",str(profile),"--report",str(report),"--smoke-test","--route","tool/workbench.photo-catalog","--theme","dark","--catalog-source",str(source),"--catalog-output",str(output),"--catalog-site",site,"--catalog-operation",operation,"--catalog-first"],env=clean_environment(),stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    memory=ProcessMemory(process.pid);latencies=[];startup_latencies=[];timeouts=0;probe_errors=[];start=time.monotonic();ready=False
    try:
        while process.poll() is None and time.monotonic()-start<55:
            if report.is_file():break  # Resultado escrito: el cierre ya no forma parte de la carga interactiva.
            handles=[]
            @callback_type
            def find(hwnd,lparam):
                pid=wintypes.DWORD();user.GetWindowThreadProcessId(hwnd,ctypes.byref(pid))
                if pid.value==process.pid and user.IsWindowVisible(hwnd):handles.append(hwnd)
                return True
            user.EnumWindows(find,0)
            if not ready:
                log=profile/"logs/workbench.jsonl"
                ready=log.is_file() and '"app.ready"' in log.read_text(encoding="utf-8",errors="replace")
            if handles:
                value=ctypes.c_size_t();t=time.perf_counter()
                ctypes.set_last_error(0)
                ok=user.SendMessageTimeoutW(handles[0],0,0,0,2,500,ctypes.byref(value));elapsed=time.perf_counter()-t
                error=ctypes.get_last_error()
                if report.is_file():break
                # Una ventana destruida no implica un bucle de mensajes bloqueado; conserva el diagnóstico.
                if not ok and error==1400:probe_errors.append({"winError":error,"secondsFromStart":round(time.monotonic()-start,3)})
                elif ready:
                    latencies.append(elapsed)
                    if not ok:timeouts+=1;probe_errors.append({"winError":error,"milliseconds":round(elapsed*1000,2),"secondsFromStart":round(time.monotonic()-start,3)})
                else:startup_latencies.append(elapsed)
            time.sleep(0.05)
        code=process.wait(timeout=5)
    except Exception:
        process.kill();process.wait(timeout=5);raise
    finally:memory_result=memory.close()
    data=json.loads(report.read_text(encoding="utf-8"));assert code==0 and data["ok"],(code,data)
    result={"ok":bool(latencies) and timeouts==0,"startupMs":data["startupMs"],"startupProbes":distribution(startup_latencies),"messagePump":distribution(latencies),"timeoutsOver500Ms":timeouts,"probeErrors":probe_errors,"memory":memory_result,"note":"Steady response starts at the app.ready marker; startup is measured separately. WM_NULL is message handling, not FPS or all interactions."}
    (root/"gui-response-validation.json").write_text(json.dumps(result,indent=2),encoding="utf-8")
    assert result["ok"],result
    return result


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,required=True);parser.add_argument("--delivery",type=Path,required=True)
    parser.add_argument("--payload",type=Path,required=True);parser.add_argument("--workdir",type=Path,required=True)
    parser.add_argument("--resume",action="store_true",help="Reuse the verified physical copy and full pre-test hashes after a test-harness failure")
    args=parser.parse_args();source=args.source.resolve();delivery=args.delivery.resolve();root=args.workdir.resolve();payload=args.payload.resolve()
    workspace=Path(__file__).resolve().parents[1]
    assert root.is_relative_to(workspace/".local/qa") and source!=delivery and not root.is_relative_to(source) and not root.is_relative_to(delivery)
    root.mkdir(parents=True,exist_ok=True)
    if args.resume:
        marker=json.loads((root/".stress-owned.json").read_text(encoding="utf-8"))
        assert marker=={"source":str(source),"delivery":str(delivery),"root":str(root)}
    elif any(root.iterdir()):raise ValueError("Use a fresh empty stress directory")
    else:(root/".stress-owned.json").write_text(json.dumps({"source":str(source),"delivery":str(delivery),"root":str(root)}),encoding="utf-8")
    started=time.perf_counter();report={"ok":False,"source":str(source),"delivery":str(delivery),"originalsReadOnly":True,"stages":{}}
    clone=root/"delivery-clone";clone.mkdir(exist_ok=args.resume);snapshot={};all_files=[];copied=0
    for label,directory in (("source",source),("delivery",delivery)):
        for path in sorted(directory.rglob("*")):
            if path.is_file():
                assert path.resolve().is_relative_to(directory) and not path.is_symlink(),str(path)
                all_files.append((label,directory,path))
    total=sum(p.stat().st_size for _,_,p in all_files);read=0;phase=time.perf_counter()
    previous_pass=[]
    if args.resume:
        snapshot=json.loads((root/"original-hashes-before.json").read_text(encoding="utf-8"))
        assert set(snapshot)=={label+"/"+path.relative_to(directory).as_posix() for label,directory,path in all_files}
        for label,directory,path in all_files:
            before=snapshot[label+"/"+path.relative_to(directory).as_posix()];stat=path.stat()
            assert before["bytes"]==stat.st_size and before["mtimeNs"]==stat.st_mtime_ns,"An original changed; do not reuse this baseline"
        log=root.with_suffix(".log")
        if log.is_file():
            for line in log.read_text(encoding="utf-8",errors="replace").splitlines():
                if line.startswith('{"phase": "all-thumbnails"'):
                    item=json.loads(line);assert item["imageErrors"]==0;previous_pass.append(item["pageMs"]/1000)
        report["baselineResumed"]=True
    for i,(label,directory,path) in enumerate(all_files):
        if args.resume:continue
        relative=path.relative_to(directory);stat=path.stat();target=None
        if label=="delivery" and (relative.parts[0]=="Entrega_Museo" or relative.as_posix() in ("estado_catalogador.json","gps_revision.json")):
            target=clone/relative;copied+=stat.st_size
        snapshot[label+"/"+relative.as_posix()]={"bytes":stat.st_size,"mtimeNs":stat.st_mtime_ns,"sha256":digest(path,target)};read+=stat.st_size
        if i%100==0 or i+1==len(all_files):print(json.dumps({"phase":"baseline-and-physical-copy","files":i+1,"totalFiles":len(all_files),"readGB":round(read/1e9,2),"totalGB":round(total/1e9,2),"elapsedSeconds":round(time.perf_counter()-phase,1)}),flush=True)
    (root/"original-hashes-before.json").write_text(json.dumps(snapshot,indent=2),encoding="utf-8")
    report["stages"]["baseline"]={"files":len(all_files),"bytesHashed":total,"physicalCopyBytes":copied if not args.resume else sum(p.stat().st_size for p in clone.rglob('*') if p.is_file()),"seconds":round(time.perf_counter()-phase,2),"resumed":args.resume}
    # Lee únicamente el libro copiado. Conserva el contexto aceptado sin inventar metadatos de campo.
    from openpyxl import load_workbook
    workbook=load_workbook(clone/"Entrega_Museo/Catalogo_fotografico_pozos.xlsx",read_only=True,data_only=True)
    headers=[c.value for c in workbook.active[1]];rows=[dict(zip(headers,[str(v or "") for v in values])) for values in workbook.active.iter_rows(min_row=2,values_only=True) if any(v is not None for v in values)];workbook.close()
    assert rows and len({r["Sitio"] for r in rows})==1
    site=rows[0]["Sitio"];operation=rows[0]["Operacion"]
    worker=Worker(payload/"modules/catalog/engine/CatalogWorker.exe");image_errors=[];items=[];page_times=[];preview_times=[]
    try:
        first,elapsed=worker.rpc("open",source=str(source),output=str(clone),site=site,operation=operation,cache=str(root/"cache"))
        assert not first.get("migrationRequired"),"Unexpected migration; the original delivery remains untouched"
        expected=first["state"]["total"];assert first["state"]["registered"]==len(rows),(first["state"],len(rows))
        report["stages"]["open"]={"seconds":elapsed,"state":first["state"],"xlsxRows":len(rows)}
        for offset in range(0,expected,60):
            page,t=worker.rpc("page",offset=offset,filter="Todas");page_times.append(t);items.extend(page["items"])
            image_errors.extend({"id":item["id"],"name":item["name"]} for item in page["items"] if not item.get("thumbnail"))
            print(json.dumps({"phase":"all-thumbnails","processed":len(items),"total":expected,"pageMs":round(t*1000,1),"imageErrors":len(image_errors)}),flush=True)
        assert len(items)==expected and len({i["id"] for i in items})==expected
        selections={0,len(items)//2,len(items)-1}
        selections.update(range(0,len(items),max(1,len(items)//24)))
        selections.update(i for i,item in enumerate(items) if Path(item["name"]).suffix.lower() in (".heic",".heif"))
        selections.add(max(range(len(items)),key=lambda i:Path(items[i]["path"]).stat().st_size))
        for index in sorted(selections):
            selected,t=worker.rpc("select",id=items[index]["id"]);preview_times.append(t)
            if selected["selection"]["preview"]["imageError"]:image_errors.append({"id":items[index]["id"],"name":items[index]["name"],"error":selected["selection"]["preview"]["imageError"]})
        cached=[]
        for offset in reversed(range(0,expected,60)):
            _,t=worker.rpc("page",offset=offset,filter="Todas");cached.append(t)
        filters={}
        for condition in ("Por catalogar","Pendientes","Omitidas","Registradas","GPS por revisar"):
            result,t=worker.rpc("page",filter=condition);filters[condition]={"matched":result["matched"],"ms":round(t*1000,2)}
        selected,_=worker.rpc("select",id=items[0]["id"]);record=selected["selection"]["record"];assert record
        form={"operation":record["Operacion"],"unitType":record["Tipo_unidad"],"unit":record["Unidad"],"level":record["Nivel"],"subunitType":record["Tipo_subunidad"],"subunit":record["Subunidad"],"view":record["Vista_o_detalle"],"notes":record["Observaciones"]}
        proposed,_=worker.rpc("propose",id=items[0]["id"],form=form);assert proposed["proposedName"]==record["Nombre_archivo"]
        corrections=[]
        for _ in range(12):
            result,t=worker.rpc("correct",id=items[0]["id"],form=form);corrections.append(t);assert result["state"]["registered"]==len(rows)
        rejected,_=worker.rpc("save",expected=False,id=items[0]["id"],form=form)
        assert rejected["errorType"]=="ValueError"
        gps,t=worker.rpc("exportGps",destination=str(root/("GPS-stress-test-"+str(time.time_ns())+".xlsx")))
        report["stages"]["photos"]={"total":expected,"allThumbnailsGenerated":len(items)-len(image_errors),"imageErrors":image_errors,"pagesCold":distribution(previous_pass or page_times),"pagesCached":distribution(cached),"previews":distribution(preview_times),"previewsMayUseCache":args.resume,"filters":filters,"sameContextCorrections":distribution(corrections),"duplicateSaveRejected":True}
        report["stages"]["gpsExport"]={"seconds":t,"result":gps["report"]}
        worker.rpc("close")
    finally:report["workerMemory"]=worker.close()
    # La decodificación representativa sin caché conserva la validez del pico de memoria
    # al reanudar, sin duplicar toda la entrega ni invalidar las mediciones de caché.
    cold=Worker(payload/"modules/catalog/engine/CatalogWorker.exe")
    try:
        _,cold_open=cold.rpc("open",source=str(source),output=str(clone),site=site,operation=operation,cache=str(root/("cold-probe-"+str(time.time_ns()))))
        heavy={max(range(len(items)),key=lambda i:Path(items[i]["path"]).stat().st_size)}
        heavy.update(i for i,item in enumerate(items) if Path(item["name"]).suffix.lower() in (".heic",".heif"))
        for index in sorted(heavy):
            item,_=cold.rpc("select",id=items[index]["id"])
            assert not item["selection"]["preview"]["imageError"],item["selection"]["preview"]
        cold.rpc("close")
    finally:report["coldProbeMemory"]=cold.close()
    report["stages"]["coldProbe"]={"initialPagePhotos":60,"additionalPreviews":len(heavy),"openSeconds":cold_open}
    (root/"stress-partial.json").write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding="utf-8")
    assert not image_errors,image_errors
    report["stages"]["gui"]=gui_probe(payload/"ArchaeoHub.exe",root,source,clone,site,operation)
    print(json.dumps({"phase":"integrity-after","files":len(all_files),"totalGB":round(total/1e9,2)}),flush=True)
    differences=[];read=0;phase=time.perf_counter()
    for i,(label,directory,path) in enumerate(all_files):
        key=label+"/"+path.relative_to(directory).as_posix();before=snapshot[key]
        stat=path.stat();actual=digest(path);read+=stat.st_size
        if actual!=before["sha256"] or stat.st_size!=before["bytes"] or stat.st_mtime_ns!=before["mtimeNs"]:differences.append(key)
        if i%200==0 or i+1==len(all_files):print(json.dumps({"phase":"integrity-after","processed":i+1,"totalFiles":len(all_files),"readGB":round(read/1e9,2)}),flush=True)
    current_count=sum(p.is_file() for directory in (source,delivery) for p in directory.rglob("*"))
    report["stages"]["integrity"]={"files":len(all_files),"bytesHashedAgain":total,"changes":differences,"fileCountUnchanged":current_count==len(all_files),"seconds":round(time.perf_counter()-phase,2)}
    assert not differences and current_count==len(all_files),differences
    report["ok"]=True;report["elapsedSeconds"]=round(time.perf_counter()-started,2)
    report["campaignElapsedSeconds"]=round(time.time()-(root/".stress-owned.json").stat().st_mtime,2)
    (root/"stress-validation.json").write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding="utf-8")
    print(json.dumps({"ok":True,"elapsedSeconds":report["elapsedSeconds"],"photos":report["stages"]["photos"]["total"],"peakWorkerMB":round(report["workerMemory"]["peakWorkingSetBytes"]/1e6,1)}),flush=True)


if __name__=="__main__":main()
