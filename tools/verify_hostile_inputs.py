"""Pruebas acotadas de entradas malformadas y regresión con fórmulas inocuas de hoja de cálculo."""
import argparse,csv,hashlib,json,os,random,subprocess,zipfile
from pathlib import Path
from openpyxl import load_workbook,Workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--payload',type=Path,required=True);parser.add_argument('--workdir',type=Path,required=True)
parser.add_argument('--baseline',action='store_true')
args=parser.parse_args();payload=args.payload.resolve();root=args.workdir.resolve();root.mkdir(parents=True,exist_ok=True)
if any(root.iterdir()):raise ValueError('Use an empty QA directory')
env={k:v for k,v in os.environ.items() if k.upper() not in ('PYTHONHOME','PYTHONPATH','QT_PLUGIN_PATH','QML2_IMPORT_PATH')}
env['PATH']=os.environ['SystemRoot']+';'+str(Path(os.environ['SystemRoot'])/'System32');env['GAMBOA_CONFIG_DIR']=str(root/'transcription')
checks=[];findings=[]

def rpc(relative,request):
 command=[str(payload/relative)]
 if relative=='CoordinateWorker.exe':command+=['--engine-dir',str(payload/'modules/coordinates/engine')]
 data=request if isinstance(request,bytes) else json.dumps(request,ensure_ascii=False).encode()
 result=subprocess.run(command,input=data,capture_output=True,timeout=15,env=env)
 assert result.returncode in (0,1,2),(relative,result.returncode,result.stderr[-400:])
 events=[json.loads(line) for line in result.stdout.splitlines()];assert events,(relative,result.stderr)
 checks.append({'worker':relative,'controlledResponse':True});return events[-1]

workers=['CoordinateWorker.exe','modules/documents/engine/DocumentWorker.exe','modules/coordinates/spreadsheets/SpreadsheetWorker.exe','modules/inventory/engine/InventoryWorker.exe']
for worker in workers:
 for data in (b'{}',b'[]',b'null',b'{',b'{"schemaVersion":999}',b'{"schemaVersion":1,"op":"not-an-operation"}'):
  reply=rpc(worker,data);assert not reply.get('ok'),(worker,reply)

rng=random.Random(1374)
bad=['NaN','Inf','1e9999','--84','', '84 60 0 W','<img src="https://example.invalid">','0'*5000,'-84,10']
for index in range(80):
 reply=rpc('CoordinateWorker.exe',{'schemaVersion':1,'op':'convert','source':'EPSG:4326','target':'EPSG:8908','inputFormat':rng.choice(['DD','DM','DMS','unknown']), 'rows':[{'id':str(index),'x':rng.choice(bad),'y':rng.choice(bad)}]})
 assert not reply['ok'] or reply['failed']==1,reply

source=root/'Malformed.pdf';source.write_bytes(b'%PDF-1.7\ntruncated\n')
assert not rpc(workers[1],{'schemaVersion':1,'paths':[str(source)]})['ok']
docx=root/'External-entity.docx'
with zipfile.ZipFile(docx,'w') as archive:
 archive.writestr('word/document.xml','<!DOCTYPE d [<!ENTITY secret SYSTEM "https://example.invalid/never-read">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:t>&secret;</w:t></w:p></w:document>')
assert not rpc(workers[1],{'schemaVersion':1,'paths':[str(docx)]})['ok']
badxlsx=root/'Malformed.xlsx';badxlsx.write_bytes(b'PK not a workbook')
assert not rpc(workers[2],{'schemaVersion':1,'op':'readSpreadsheet','path':str(badxlsx)})['ok']

# La fórmula inocua =1+1 detecta su interpretación sin abrir Excel ni ejecutar comandos.
request=root/'coordinates.json';response=root/'coordinates-result.json';output=root/'Coordinates.csv'
request.write_text(json.dumps({'op':'convert','source':'EPSG:4326','target':'EPSG:4326','headers':['=1+1','x','y'],'rows':[{'id':'=1+1','x':'-84','y':'10','fields':['=1+1','-84','10']}]}),encoding='utf-8')
process=subprocess.run([str(payload/'ArchaeoHub.exe'),'--data-dir',str(root/'coordinate-profile'),'--coordinate-request',str(request),'--coordinate-export',str(output),'--report',str(response)],env=env,capture_output=True,timeout=30)
assert process.returncode==0 and json.loads(response.read_text())['rows'][0]['id']=='=1+1'
with output.open(encoding='utf-8-sig',newline='') as stream:cells=list(csv.reader(stream))
if cells[0][0]=='=1+1' or cells[1][0]=='=1+1':findings.append('coordinate-csv-formula')

db=root/'Inventory.ahinv';worker=workers[3]
def inventory(op,**kwargs):
 reply=rpc(worker,{'schemaVersion':1,'database':str(db),'op':op,**kwargs});assert reply['ok'],reply;return reply['result']
inventory('create',metadata={'project':'Synthetic QA','site':'QA'})
inventory('add',record={'operacion':'2','q_cer':1,'=1+1':'literal','comentarios':'=1+1'})
for suffix in ('csv','xlsx'):
 output=root/('Inventory.'+suffix);inventory('export',output=str(output))
 if suffix=='csv':
  with output.open(encoding='utf-8-sig',newline='') as stream:header=next(csv.reader(stream))
  if '=1+1' in header:findings.append('inventory-csv-header-formula')
 else:
  book=load_workbook(output,data_only=False)
  if any(cell.data_type=='f' for sheet in book for row in sheet for cell in row):findings.append('inventory-xlsx-header-formula')
  book.close()
assert '=1+1' in inventory('open')['rows'][0]
keep=hashlib.sha256(db.read_bytes()).hexdigest()
for invalid in (-1,float('nan'),1.5,True,'=1+1'):
 reply=rpc(worker,{'schemaVersion':1,'database':str(db),'op':'add','record':{'operacion':'2','q_cer':invalid}})
 assert not reply['ok']
assert hashlib.sha256(db.read_bytes()).hexdigest()==keep
reply=rpc(worker,{'schemaVersion':1,'database':str(db),'op':'export','output':str(root/'Inventory.xlsx')});assert not reply['ok']

if not args.baseline:
 # Un ZIP pequeño que declara una hoja enorme debe rechazarse sin recorrer sus celdas.
 folder=root/'Catalog'/'Entrega_Museo';folder.mkdir(parents=True);path=folder/'Catalogo_fotografico_pozos.xlsx'
 book=Workbook();book.active.append(['test']);book.save(path);book.close()
 import re
 with zipfile.ZipFile(path) as archive:entries={name:archive.read(name) for name in archive.namelist()}
 entries['xl/worksheets/sheet1.xml']=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:XFD1048576"',entries['xl/worksheets/sheet1.xml'])
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
  for name,data in entries.items():archive.writestr(name,data)
 reply=rpc('modules/catalog/engine/CatalogWorker.exe',(json.dumps({'schemaVersion':1,'op':'load','location':str(path)})+'\n').encode())
 assert not reply['ok'] and '100.000' in reply['error'],reply
 assert not findings,findings
result={'ok':not findings,'baseline':args.baseline,'controlledResponses':len(checks),'findings':findings,'inventoryUnchangedAfterInvalidWrites':True,'rawFormulaTextPreservedInDatabaseAndJson':True,'hostSha256':hashlib.sha256((payload/'ArchaeoHub.exe').read_bytes()).hexdigest()}
(root/'hostile-inputs-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result),flush=True)
