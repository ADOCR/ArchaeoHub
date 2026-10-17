"""Acota los libros de catálogo externos antes de recorrer hojas o migrarlos."""
from pathlib import Path
import zipfile
from openpyxl import load_workbook as _load_workbook


def load_workbook(filename, *args, **kwargs):
    path = Path(filename)
    if not path.is_file() or path.stat().st_size > 64 * 1024 * 1024:
        raise ValueError("El Excel del catálogo supera 64 MB o no está disponible.")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
            raise ValueError("El contenido expandido del catálogo supera 128 MB o 4.096 entradas.")
    kwargs["keep_links"] = False
    book = _load_workbook(filename, *args, **kwargs)
    if any((sheet.max_row or 0) > 100001 or (sheet.max_column or 0) > 128 for sheet in book):
        book.close()
        raise ValueError("El catálogo supera 100.000 registros o 128 columnas. Divide el catálogo antes de abrirlo.")
    return book
