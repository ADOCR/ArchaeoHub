"""Reporte local de referencias GPS de fotografías catalogadas.

Las coordenadas proceden del EXIF de las copias. No son mediciones de la
posición exacta de la unidad arqueológica ni se añaden al Excel del Museo.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import statistics
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:
    pass


GREEN = "2F5D50"
PALE_GREEN = "E8F0ED"
PALE_AMBER = "FFF1CF"
RED = "A94A45"
DISPERSION_ALERT_METERS = 25.0  # Alerta de revisión, no precisión arqueológica.
REQUIRED_COLUMNS = (
    "Sitio", "Operacion", "Tipo_unidad", "Unidad", "Tipo_subunidad",
    "Subunidad", "Nivel", "Nombre_archivo", "Nombre_original",
)
UNIT_HEADERS = (
    "Sitio", "Operación", "Tipo de unidad", "Número / código",
    "Unidad en catálogo", "Tipo de subunidad", "Subunidad",
    "Latitud GPS de referencia (°)", "Longitud GPS de referencia (°)",
    "Altitud EXIF (m)", "Fotos con GPS válido", "Fotos catalogadas",
    "Separación máxima (m)", "Foto de referencia", "Estado",
    "Latitud unidad validada (°)", "Longitud unidad validada (°)",
)
PHOTO_HEADERS = (
    "Sitio", "Operación", "Tipo de unidad", "Número / código",
    "Tipo de subunidad", "Subunidad", "Nivel", "Archivo catalogado",
    "Archivo original", "Latitud GPS foto (°)", "Longitud GPS foto (°)",
    "Altitud EXIF (m)", "Error horizontal EXIF (m)", "Datum EXIF",
    "Estado GPS",
)


@dataclass(frozen=True)
class GPSPoint:
    latitude: float
    longitude: float
    altitude: float | None = None
    horizontal_error: float | None = None
    datum: str = ""


@dataclass(frozen=True)
class ReportResult:
    path: Path
    units: int
    photos: int
    photos_with_gps: int
    units_without_gps: int
    photos_invalid_gps: int
    photos_without_gps: int
    photos_unreadable: int


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _text(value: object) -> str:
    """Mantiene identificadores como texto sin permitir fórmulas en el reporte."""
    result = str(value or "")
    if result.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + result
    return result


def _tag_text(value: object) -> str:
    if isinstance(value, bytes):
        return value.decode("ascii", errors="replace").strip("\x00 ")
    return str(value or "").strip("\x00 ")


def _dms_to_decimal(values: object, reference: object, positive: str, negative: str) -> float:
    if not isinstance(values, (tuple, list)) or len(values) != 3:
        raise ValueError("Faltan grados, minutos o segundos GPS.")
    direction = _tag_text(reference).upper()
    if direction not in {positive, negative}:
        raise ValueError("Falta la referencia cardinal GPS.")
    degrees, minutes, seconds = (float(part) for part in values)
    if not all(math.isfinite(part) and part >= 0 for part in (degrees, minutes, seconds)):
        raise ValueError("Coordenada GPS no numérica.")
    if minutes >= 60 or seconds >= 60:
        raise ValueError("Minutos o segundos GPS fuera de rango.")
    decimal = degrees + minutes / 60 + seconds / 3600
    return -decimal if direction == negative else decimal


def read_gps(path: Path) -> tuple[GPSPoint | None, str]:
    """Lee el GPS EXIF sin cargar ni modificar los píxeles de la fotografía."""
    try:
        with Image.open(path) as image:
            gps = image.getexif().get_ifd(0x8825)
    except Exception:
        return None, "No se pudo leer la copia"
    if not gps:
        return None, "Sin GPS EXIF"
    try:
        latitude = _dms_to_decimal(gps.get(2), gps.get(1), "N", "S")
        longitude = _dms_to_decimal(gps.get(4), gps.get(3), "E", "W")
        if abs(latitude) > 90 or abs(longitude) > 180:
            raise ValueError("Coordenada fuera de rango.")
        altitude = None
        if gps.get(6) is not None:
            try:
                candidate = float(gps[6])
                if math.isfinite(candidate):
                    altitude_ref = gps.get(5, 0)
                    if isinstance(altitude_ref, bytes):
                        altitude_ref = altitude_ref[0] if altitude_ref else 0
                    altitude = -candidate if altitude_ref == 1 else candidate
            except (TypeError, ValueError, ZeroDivisionError, OverflowError):
                pass
        error = None
        if gps.get(31) is not None:
            try:
                candidate = float(gps[31])
                if math.isfinite(candidate) and candidate >= 0:
                    error = candidate
            except (TypeError, ValueError, ZeroDivisionError, OverflowError):
                pass
        return GPSPoint(latitude, longitude, altitude, error, _tag_text(gps.get(18))), "Con GPS EXIF"
    except (TypeError, ValueError, ZeroDivisionError, OverflowError):
        return None, "GPS EXIF no válido"


def _distance_m(first: GPSPoint, second: GPSPoint) -> float:
    """Distancia aproximada sobre WGS 84 solo para comparar dispersión."""
    lat1, lon1 = math.radians(first.latitude), math.radians(first.longitude)
    lat2, lon2 = math.radians(second.latitude), math.radians(second.longitude)
    a = math.sin((lat2 - lat1) / 2) ** 2
    a += math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 2 * 6371008.8 * math.asin(min(1.0, math.sqrt(a)))


def _unit_number(unit_type: str, unit: str) -> str:
    upper = unit.upper()
    if unit_type == "Trinchera":
        if upper.startswith("TR") and len(unit) > 2:
            return unit[2:]
        if upper.startswith("T") and len(unit) > 1:
            return unit[1:]
    if unit_type == "Pozo auxiliar":
        if upper.startswith("PA") and len(unit) > 2:
            return unit[2:]
        if upper.startswith("A") and len(unit) > 1:
            return unit[1:]
    return unit


def _catalog_rows(catalog_xlsx: Path, state_path: Path) -> tuple[list[dict[str, object]], str, str]:
    if not catalog_xlsx.is_file() or not state_path.is_file():
        raise FileNotFoundError("Se necesita el Excel y el estado del catálogo para generar el reporte.")
    if (state_path.parent / ".transaccion_catalogador.json").exists():
        raise RuntimeError("Hay un guardado pendiente. Reabre el catálogo antes de exportar coordenadas.")
    excel_bytes = catalog_xlsx.read_bytes()
    state_bytes = state_path.read_bytes()
    workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        headers = next(iterator)
        if not set(REQUIRED_COLUMNS).issubset(headers):
            raise ValueError("El Excel del catálogo no tiene las columnas necesarias.")
        rows = [
            dict(zip(headers, values))
            for values in iterator
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()
    state = json.loads(state_bytes)
    if len(set(state.get("processed_paths", []))) != len(rows):
        raise ValueError(
            "El estado y el Excel no coinciden. No se generó el reporte; "
            "revisa el catálogo antes de continuar."
        )
    filenames = [_text(row.get("Nombre_archivo")) for row in rows]
    if len(filenames) != len(set(filenames)):
        raise ValueError("Hay nombres de fotografía repetidos en el Excel.")
    return rows, _sha256(excel_bytes), _sha256(state_bytes)


def _style_sheet(sheet, widths: tuple[int, ...], validated_columns: tuple[int, ...] = ()) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="24312D")
            cell.alignment = Alignment(vertical="center")
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F5F8F6")
        for column in validated_columns:
            row[column - 1].fill = PatternFill("solid", fgColor=PALE_AMBER)
        for cell in row:
            if isinstance(cell.value, str) and (
                cell.value.startswith("Sin GPS") or cell.value.startswith("GPS EXIF no válido")
                or cell.value.startswith("No se pudo") or cell.value.startswith("Revisar dispersión")
            ):
                cell.font = Font(name="Arial", size=10, color=RED)
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def _workbook_for_rows(rows: list[dict[str, object]], photos_dir: Path) -> tuple[Workbook, ReportResult]:
    groups: dict[tuple[str, ...], list[tuple[dict[str, object], GPSPoint | None]]] = defaultdict(list)
    details: list[tuple[object, ...]] = []
    with_gps = 0
    invalid_gps = 0
    without_gps = 0
    unreadable = 0
    for row in rows:
        site = _text(row.get("Sitio"))
        operation = _text(row.get("Operacion"))
        unit_type = _text(row.get("Tipo_unidad"))
        unit = _text(row.get("Unidad"))
        subunit_type = _text(row.get("Tipo_subunidad"))
        subunit = _text(row.get("Subunidad"))
        filename = _text(row.get("Nombre_archivo"))
        if not filename or Path(filename).name != filename or "/" in filename or "\\" in filename:
            raise ValueError(f"Nombre de copia no válido en el Excel: {filename!r}")
        copy = photos_dir / filename
        if copy.is_file():
            point, status = read_gps(copy)
        else:
            point, status = None, "Copia ausente"
        if point is not None:
            with_gps += 1
        elif status == "GPS EXIF no válido":
            invalid_gps += 1
        elif status == "Sin GPS EXIF":
            without_gps += 1
        else:
            unreadable += 1
        key = (site, operation, unit_type, unit, subunit_type, subunit)
        groups[key].append((row, point))
        details.append((
            site, operation, unit_type, _unit_number(unit_type, unit),
            subunit_type, subunit, _text(row.get("Nivel")),
            filename, _text(row.get("Nombre_original")),
            point.latitude if point else None,
            point.longitude if point else None,
            point.altitude if point else None,
            point.horizontal_error if point else None,
            _text(point.datum) if point else "",
            status,
        ))

    workbook = Workbook()
    units_sheet = workbook.active
    units_sheet.title = "Unidades"
    units_sheet.append(UNIT_HEADERS)
    units_without_gps = 0
    for key in sorted(groups, key=lambda item: tuple(part.casefold() for part in item)):
        site, operation, unit_type, unit, subunit_type, subunit = key
        entries = groups[key]
        points = [(row, point) for row, point in entries if point is not None]
        if points:
            center = GPSPoint(
                statistics.median(point.latitude for _, point in points),
                statistics.median(point.longitude for _, point in points),
            )
            reference_row, reference = min(
                points,
                key=lambda item: (
                    _distance_m(item[1], center),
                    _text(item[0].get("Nombre_archivo")),
                ),
            )
            maximum = max(_distance_m(reference, point) for _, point in points)
            status = (
                "Revisar dispersión GPS; validar unidad"
                if maximum > DISPERSION_ALERT_METERS
                else "GPS de foto; validar unidad"
            )
            latitude, longitude, altitude = reference.latitude, reference.longitude, reference.altitude
            reference_name = _text(reference_row.get("Nombre_archivo"))
        else:
            units_without_gps += 1
            latitude = longitude = altitude = maximum = None
            reference_name = ""
            status = "Sin GPS EXIF; ubicar unidad"
        units_sheet.append((
            site, operation, unit_type, _unit_number(unit_type, unit),
            unit, subunit_type, subunit, latitude, longitude, altitude,
            len(points), len(entries), round(maximum, 1) if maximum is not None else None,
            reference_name, status, None, None,
        ))

    photos_sheet = workbook.create_sheet("Fotos GPS")
    photos_sheet.append(PHOTO_HEADERS)
    for detail in details:
        photos_sheet.append(detail)
    _style_sheet(
        units_sheet,
        (16, 13, 19, 19, 20, 20, 18, 21, 22, 17, 15, 18, 21, 42, 39, 23, 24),
        validated_columns=(16, 17),
    )
    _style_sheet(
        photos_sheet,
        (16, 13, 19, 19, 20, 18, 13, 43, 32, 20, 21, 17, 22, 20, 26),
    )
    for sheet, columns in ((units_sheet, (8, 9, 16, 17)), (photos_sheet, (10, 11))):
        for column in columns:
            for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for cell in cells:
                    cell.number_format = "0.00000000"
    for sheet, columns in ((units_sheet, (10, 13)), (photos_sheet, (12, 13))):
        for column in columns:
            for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for cell in cells:
                    cell.number_format = "0.0"
    return workbook, ReportResult(
        Path(), len(groups), len(rows), with_gps, units_without_gps,
        invalid_gps, without_gps, unreadable,
    )


def export_gps_report(
    catalog_xlsx: Path,
    photos_dir: Path,
    state_path: Path,
    destination: Path,
) -> ReportResult:
    """Genera un Excel independiente sin modificar el catálogo ni las fotos."""
    catalog_xlsx = Path(catalog_xlsx)
    photos_dir = Path(photos_dir)
    state_path = Path(state_path)
    destination = Path(destination)
    delivery = catalog_xlsx.parent.resolve()
    if destination.resolve().is_relative_to(delivery):
        raise ValueError("Guarda las coordenadas fuera de Entrega_Museo.")
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("El reporte debe tener extensión .xlsx.")
    if destination.exists():
        raise FileExistsError(f"El reporte ya existe: {destination.name}")
    rows, excel_hash, state_hash = _catalog_rows(catalog_xlsx, state_path)
    workbook, result = _workbook_for_rows(rows, photos_dir)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix="coordenadas_", suffix=".xlsx", dir=destination.parent
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        if _sha256(catalog_xlsx.read_bytes()) != excel_hash:
            raise RuntimeError("El Excel cambió durante la exportación; inténtalo otra vez.")
        if _sha256(state_path.read_bytes()) != state_hash:
            raise RuntimeError("El estado cambió durante la exportación; inténtalo otra vez.")
        if destination.exists():
            raise FileExistsError(f"El reporte ya existe: {destination.name}")
        os.replace(temp_path, destination)
    finally:
        temp_path.unlink(missing_ok=True)
    return ReportResult(
        destination, result.units, result.photos, result.photos_with_gps,
        result.units_without_gps, result.photos_invalid_gps,
        result.photos_without_gps, result.photos_unreadable,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera referencias GPS desde un catálogo existente.")
    parser.add_argument("carpeta_salida", type=Path, help="Carpeta de trabajo que contiene Entrega_Museo")
    parser.add_argument("reporte", type=Path, help="Nuevo archivo .xlsx fuera de Entrega_Museo")
    args = parser.parse_args()
    output = args.carpeta_salida
    result = export_gps_report(
        output / "Entrega_Museo" / "Catalogo_fotografico_pozos.xlsx",
        output / "Entrega_Museo" / "Fotos",
        output / "estado_catalogador.json",
        args.reporte,
    )
    print(
        f"{result.path}: {result.units} unidades; "
        f"{result.photos_with_gps}/{result.photos} fotografías con GPS válido; "
        f"{result.photos_invalid_gps} con GPS inválido; "
        f"{result.photos_without_gps} sin GPS."
    )


if __name__ == "__main__":
    main()
