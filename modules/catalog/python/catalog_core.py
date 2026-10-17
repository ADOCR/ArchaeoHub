"""Catalogador local de fotografías arqueológicas para Windows.

Los originales nunca se modifican. La rotación, el zoom y el desplazamiento
pertenecen únicamente al visor.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import queue
import re
import shutil
import tempfile
import threading
from collections import OrderedDict, deque
from datetime import datetime
from pathlib import Path
from typing import Callable

from PIL import Image, ImageDraw, ImageOps
from openpyxl import Workbook
from catalog_workbook import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from georreferenciacion import export_gps_report, read_gps
from gps_manual import add_manual_gps

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
    HEIC_ENABLED = True
except ImportError:
    HEIC_ENABLED = False


APP_DIR = Path(__file__).resolve().parent
ASSET_PATH = APP_DIR / "assets" / "cucharilla.png"
STATE_VERSION = 5
SUFFIXES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"}
if HEIC_ENABLED:
    SUFFIXES.update({".heic", ".heif"})

FIELDS = [
    "Sitio", "Operacion", "Tipo_unidad", "Unidad", "Tipo_subunidad",
    "Subunidad", "Nivel",
    "Foto_en_unidad_y_nivel", "Nombre_archivo", "Nombre_original",
    "Fecha_EXIF", "Vista_o_detalle", "Observaciones", "Fecha_catalogacion",
]
FIELDS_V4 = [
    "Sitio", "Operacion", "Tipo_unidad", "Unidad", "Nivel",
    "Foto_en_unidad_y_nivel", "Nombre_archivo", "Nombre_original",
    "Fecha_EXIF", "Vista_o_detalle", "Observaciones", "Fecha_catalogacion",
]
LEGACY_FIELDS_V3 = [
    "Sitio", "Pozo", "Nivel", "Foto_en_pozo_y_nivel", "Nombre_archivo",
    "Nombre_original", "Fecha_EXIF", "Vista_o_detalle", "Observaciones",
    "Fecha_catalogacion",
]
OLD_FIELDS = [
    "Sitio", "Pozo", "Nivel", "Foto_en_pozo_y_nivel", "Nombre_nuevo",
    "Ruta_relativa", "Nombre_original", "Ruta_original", "Fecha_EXIF",
    "Vista_o_detalle", "Observaciones", "Fecha_catalogacion",
]

SAND = "#F5F2EB"
WHITE = "#FFFFFF"
INK = "#24312D"
MUTED = "#65716C"
GREEN = "#2F5D50"
GREEN_DARK = "#24493F"
GREEN_PALE = "#E8F0ED"
GOLD = "#B8833B"
RED = "#A94A45"
CANVAS_BG = "#202624"
ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
UNIT_TYPES = ("Pozo", "Pozo auxiliar", "Trinchera")
PHOTO_FILTERS = (
    "Todas", "Por catalogar", "Pendientes", "Omitidas",
    "Registradas", "GPS por revisar",
)
SUBUNIT_TYPES = (
    "", "Cuadro", "Cuadrante", "Suboperación", "Ampliación",
    "Rasgo cultural", "Rasgo funerario", "Otro",
)
SUBUNIT_FILENAME_CODES = {
    "Cuadro": "CD",
    "Cuadrante": "CUAD",
    "Suboperación": "SUBOP",
    "Ampliación": "AMP",
    "Rasgo cultural": "RC",
    "Rasgo funerario": "RF",
    "Otro": "SU",
}


def enable_windows_dpi_awareness() -> None:
    """Permite que Tk reciba dimensiones correctas con escalado de Windows."""
    if os.name != "nt":
        return
    try:
        import ctypes

        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def excel_safe(value: object) -> str:
    """Convierte a texto válido para XLSX y evita fórmulas accidentales."""
    # Algunas cámaras escriben fechas EXIF rellenadas con NUL. Esos controles
    # no son válidos en XML/XLSX, aunque el archivo de imagen sí pueda abrirse.
    text = ILLEGAL_EXCEL_CHARACTERS.sub("", str(value or ""))
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def normalized(path: Path) -> str:
    return str(path.resolve())


def discover(source: Path, output: Path) -> list[Path]:
    files: list[Path] = []
    output_resolved = output.resolve()
    for path in source.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUFFIXES:
            continue
        try:
            if path.resolve().is_relative_to(output_resolved):
                continue
        except (OSError, ValueError):
            continue
        files.append(path)
    return sorted(files, key=lambda item: (str(item.parent).casefold(), item.name.casefold()))


class SaveCommittedWarning(Exception):
    """La foto y el Excel se guardaron, pero el estado requiere recuperación."""


class CorrectionCommittedWarning(Exception):
    """La corrección se aplicó, pero el estado requiere recuperación."""


class WithdrawalCommittedWarning(Exception):
    """El registro se retiró, pero falta completar un archivo auxiliar."""


class DiscardCommittedWarning(Exception):
    """El original quedó resguardado; al reabrir se completará el estado."""


class CatalogLock:
    """Impide abrir dos ventanas nuevas sobre la misma carpeta de trabajo."""

    def __init__(self, output: Path):
        self.file = (output / ".catalogador.lock").open("a+b")
        try:
            self.file.seek(0, os.SEEK_END)
            if self.file.tell() == 0:
                self.file.write(b"1")
                self.file.flush()
            self.file.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(self.file.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.file.close()
            raise ValueError(
                "Esta carpeta de trabajo ya está abierta en otra ventana del "
                "catalogador. Ciérrala antes de continuar."
            ) from exc

    def close(self) -> None:
        if self.file.closed:
            return
        self.file.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(self.file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.file.fileno(), fcntl.LOCK_UN)
        finally:
            self.file.close()


class Catalog:
    """Persistencia del catálogo y transacciones de copia/Excel/estado."""

    def __init__(
        self, source: Path, output: Path, site: str, legacy_operation: str | None = None
    ):
        self.source = source.resolve()
        self.output = output.resolve()
        self.site = site
        self.legacy_operation = legacy_operation
        self.delivery = self.output / "Entrega_Museo"
        self.photos = self.delivery / "Fotos"
        self.xlsx_path = self.delivery / "Catalogo_fotografico_pozos.xlsx"
        self.state_path = self.output / "estado_catalogador.json"
        self.journal_path = self.output / ".transaccion_catalogador.json"
        self.gps_review_path = self.output / "gps_revision.json"
        self.output.mkdir(parents=True, exist_ok=True)
        self.photos.mkdir(parents=True, exist_ok=True)
        self.rows: list[dict[str, str]] = []
        self.completed: set[str] = set()
        self.skipped: set[str] = set()
        self.record_map: dict[str, str] = {}
        self.corrections: list[dict[str, object]] = []
        self.gps_review: set[str] = set()
        self.gps_manual_history: list[dict[str, object]] = []
        self.last_path = ""
        self.needs_schema_migration = False
        self.needs_subunit_migration = False
        self._state_digest: bytes | None = None

        if not self.xlsx_path.exists() and not self.state_path.exists():
            self._migrate_previous()
        self._open_workbook()
        self._load_state()
        self._recover_interrupted_transaction()
        if self.needs_schema_migration:
            self._migrate_operation_and_units(legacy_operation)
        elif self.needs_subunit_migration:
            self._migrate_optional_subunits()
        self.files = discover(self.source, self.output)
        self._rebuild_record_map()
        self._load_gps_review()

    def _open_workbook(self) -> None:
        if self.xlsx_path.exists():
            self.wb = load_workbook(self.xlsx_path)
            self.ws = self.wb.active
            headers = [cell.value for cell in self.ws[1]]
            if headers == FIELDS:
                row_fields = FIELDS
            elif headers == FIELDS_V4:
                row_fields = FIELDS_V4
                self.needs_subunit_migration = True
            elif headers == LEGACY_FIELDS_V3:
                row_fields = LEGACY_FIELDS_V3
                self.needs_schema_migration = True
            else:
                raise ValueError("El Excel existente tiene columnas diferentes.")
            self.rows = [
                dict(zip(row_fields, [str(value or "") for value in values]))
                for values in self.ws.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in values)
            ]
            return

        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Fotografías"
        self.ws.append(FIELDS)
        self.ws.freeze_panes = "H2"
        self.ws.sheet_view.showGridLines = False
        self.ws.row_dimensions[1].height = 31
        for cell in self.ws[1]:
            cell.fill = PatternFill("solid", fgColor="2F5D50")
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", horizontal="center")
        widths = [14, 12, 17, 15, 18, 16, 12, 24, 48, 35, 23, 34, 52, 22]
        for index, width in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(index)].width = width
        self._write_workbook()

    def _load_state(self) -> None:
        if self.state_path.exists():
            try:
                raw_state = self.state_path.read_bytes()
                state = json.loads(raw_state.decode("utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise ValueError(f"No se pudo leer el estado del catálogo: {exc}") from exc
            self._state_digest = hashlib.sha256(raw_state).digest()
            # Compatibilidad: la versión anterior solo tenía processed_paths.
            self.completed = set(state.get("processed_paths", []))
            self.skipped = set(state.get("skipped_paths", [])) - self.completed
            self.last_path = str(state.get("last_path", ""))
            self.record_map = {
                str(path): str(filename)
                for path, filename in dict(state.get("record_map", {})).items()
            }
            self.corrections = list(state.get("corrections", []))
            if len(self.completed) != len(self.rows) and not self.journal_path.exists():
                raise ValueError(
                    "El estado y el Excel no coinciden. Puede haber quedado abierta "
                    "otra ventana con datos anteriores. Cierra todas las ventanas y "
                    "respalda la carpeta de trabajo antes de corregirlos."
                )
        elif self.rows and not self.journal_path.exists():
            raise ValueError(
                "Falta estado_catalogador.json; se necesita para reanudar sin repetir fotos."
            )

    def _state_data(
        self,
        completed: set[str] | None = None,
        skipped: set[str] | None = None,
        last_path: str | None = None,
        record_map: dict[str, str] | None = None,
        corrections: list[dict[str, object]] | None = None,
    ) -> dict[str, object]:
        return {
            "version": STATE_VERSION,
            "processed_paths": sorted(self.completed if completed is None else completed),
            "skipped_paths": sorted(self.skipped if skipped is None else skipped),
            "last_path": self.last_path if last_path is None else last_path,
            "record_map": dict(sorted(
                (self.record_map if record_map is None else record_map).items()
            )),
            "corrections": self.corrections if corrections is None else corrections,
        }

    def _atomic_json(self, path: Path, data: dict[str, object]) -> None:
        fd, temp_name = tempfile.mkstemp(prefix=path.stem + "_", suffix=".tmp", dir=path.parent)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            temp_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temp_path, path)
        finally:
            temp_path.unlink(missing_ok=True)

    def _assert_state_current(self) -> None:
        """Una instancia antigua no debe volver a escribir un estado ya cambiado."""
        try:
            current = (hashlib.sha256(self.state_path.read_bytes()).digest()
                       if self.state_path.exists() else None)
        except OSError as exc:
            raise OSError("No se pudo comprobar el estado actual del catálogo.") from exc
        if current != self._state_digest:
            raise OSError(
                "El catálogo cambió en otra ventana desde que lo abriste. No se guardó "
                "nada; cierra esta ventana y vuelve a cargar el catálogo."
            )

    def _load_gps_review(self) -> None:
        """La lista GPS vive fuera de la entrega y del estado histórico."""
        if not self.gps_review_path.exists():
            return
        try:
            data = json.loads(self.gps_review_path.read_text(encoding="utf-8"))
            self.gps_review = set(map(str, data.get("pending_paths", [])))
            self.gps_manual_history = list(data.get("manual_entries", []))
        except (OSError, ValueError, TypeError) as exc:
            raise ValueError(f"No se pudo leer la lista GPS por revisar: {exc}") from exc

    def _write_gps_review(self, pending: set[str], history: list[dict[str, object]]) -> None:
        self._atomic_json(self.gps_review_path, {
            "version": 1, "pending_paths": sorted(pending), "manual_entries": history,
        })

    def needs_gps_review(self, path: Path) -> bool:
        return normalized(path) in self.gps_review

    def gps_review_files(self) -> list[Path]:
        return [path for path in self.files if self.needs_gps_review(path)]

    def display_status(self, path: Path) -> str:
        status = self.status_of(path)
        return "gps_review" if status == "registered" and self.needs_gps_review(path) else status

    @staticmethod
    def _replace_copy_bytes(destination: Path, contents: bytes) -> None:
        fd, temp_name = tempfile.mkstemp(
            prefix="gps_copia_", suffix=destination.suffix, dir=destination.parent
        )
        temp_path = Path(temp_name)
        try:
            with os.fdopen(fd, "wb") as stream:
                stream.write(contents)
                stream.flush()
                os.fsync(stream.fileno())
            shutil.copystat(destination, temp_path)
            os.replace(temp_path, destination)
        finally:
            temp_path.unlink(missing_ok=True)

    def update_manual_gps(
        self, path: Path, latitude: str, longitude: str,
        altitude: str = "", accuracy: str = "",
    ) -> Path:
        """Corrige una copia ya registrada; nunca modifica el original ni Excel."""
        original = normalized(path)
        if original not in self.completed:
            raise ValueError("La fotografía debe estar registrada para corregir su GPS.")
        filename = self.record_map.get(original)
        if not filename or Path(filename).name != filename:
            raise ValueError("No se encontró inequívocamente la copia catalogada.")
        destination = self.photos / filename
        if destination.suffix.lower() not in {".jpg", ".jpeg"}:
            raise ValueError("El GPS manual sin recomprimir solo está disponible para JPEG.")
        if not destination.is_file():
            raise FileNotFoundError(f"No existe la copia {filename}.")
        if read_gps(destination)[0] is not None:
            raise ValueError("Esta copia ya contiene GPS válido; no se sobrescribió.")
        old_bytes = destination.read_bytes()
        new_bytes = add_manual_gps(old_bytes, latitude, longitude, altitude, accuracy)
        backup_dir = self.output / "Respaldos_GPS_manual"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{filename}"
        if backup.exists():
            raise FileExistsError("Ya existe el nombre del respaldo GPS.")
        shutil.copy2(destination, backup)
        if backup.read_bytes() != old_bytes:
            raise RuntimeError("El respaldo GPS no coincide con la copia; no se modificó.")
        pending = self.gps_review - {original}
        history = [*self.gps_manual_history, {
            "original": original, "filename": filename,
            "latitude": latitude.strip(), "longitude": longitude.strip(),
            "altitude_m": altitude.strip(), "accuracy_m": accuracy.strip(),
            "source": "transcripción manual de coordenadas de cámara",
            "recorded_at": datetime.now().isoformat(timespec="seconds"),
            "backup": str(backup),
        }]
        try:
            self._replace_copy_bytes(destination, new_bytes)
            if read_gps(destination)[0] is None:
                raise RuntimeError("No se pudo comprobar el GPS escrito en la copia.")
            self._write_gps_review(pending, history)
        except Exception:
            self._replace_copy_bytes(destination, old_bytes)
            raise
        self.gps_review, self.gps_manual_history = pending, history
        return destination

    def _write_state(
        self,
        completed: set[str] | None = None,
        skipped: set[str] | None = None,
        last_path: str | None = None,
        record_map: dict[str, str] | None = None,
        corrections: list[dict[str, object]] | None = None,
    ) -> None:
        self._assert_state_current()
        self._atomic_json(
            self.state_path,
            self._state_data(completed, skipped, last_path, record_map, corrections),
        )
        self._state_digest = hashlib.sha256(self.state_path.read_bytes()).digest()

    def remember_current(self, path: Path | None) -> None:
        value = normalized(path) if path else ""
        self._write_state(last_path=value)
        self.last_path = value

    def _write_journal(self, data: dict[str, object]) -> None:
        self._assert_state_current()
        self._atomic_json(self.journal_path, data)

    def _recover_interrupted_transaction(self) -> None:
        """Finaliza o revierte una operación cortada entre Excel y estado."""
        if not self.journal_path.exists():
            return
        try:
            journal = json.loads(self.journal_path.read_text(encoding="utf-8"))
            original = str(journal["original"])
            operation = str(journal.get("operation", "save"))
            filename = str(journal.get("filename", journal.get("new_filename", "")))
        except Exception as exc:
            raise ValueError(
                "Existe una transacción incompleta que no se puede interpretar: "
                f"{self.journal_path.name}. {exc}"
            ) from exc

        if operation == "correction":
            self._recover_interrupted_correction(journal)
            return
        if operation == "withdraw":
            self._recover_interrupted_withdrawal(journal)
            return
        if operation == "discard_unregistered":
            self._recover_interrupted_discard(journal)
            return

        destination = self.photos / filename
        excel_has_row = any(row.get("Nombre_archivo") == filename for row in self.rows)
        if excel_has_row and destination.is_file():
            completed = self.completed | {original}
            skipped = self.skipped - {original}
            record_map = dict(self.record_map)
            record_map[original] = filename
            self._write_state(
                completed=completed, skipped=skipped, last_path=original,
                record_map=record_map,
            )
            self.completed, self.skipped, self.last_path = completed, skipped, original
            self.record_map = record_map
            self.journal_path.unlink(missing_ok=True)
            return

        if destination.exists() and not excel_has_row:
            destination.unlink()
        self.journal_path.unlink(missing_ok=True)

    def _recover_interrupted_withdrawal(self, journal: dict[str, object]) -> None:
        """Restaura la copia o termina la baja según lo que confirmó Excel."""
        original = str(journal["original"])
        filename = str(journal["filename"])
        backup = Path(str(journal["backup"]))
        destination = self.photos / filename
        excel_has_row = any(row["Nombre_archivo"] == filename for row in self.rows)
        if excel_has_row:
            if backup.exists() and not destination.exists():
                os.replace(backup, destination)
            if not destination.is_file():
                raise ValueError("La baja interrumpida no tiene copia recuperable.")
            self.journal_path.unlink(missing_ok=True)
            return
        if not backup.is_file() or destination.exists():
            raise ValueError("La baja interrumpida requiere revisión manual de la copia.")
        completed = self.completed - {original}
        skipped = self.skipped | {original}
        record_map = dict(self.record_map)
        record_map.pop(original, None)
        self._write_state(
            completed=completed, skipped=skipped, last_path=original,
            record_map=record_map,
        )
        self.completed, self.skipped, self.last_path = completed, skipped, original
        self.record_map = record_map
        if self.gps_review_path.exists():
            review = json.loads(self.gps_review_path.read_text(encoding="utf-8"))
            pending = set(map(str, review.get("pending_paths", []))) - {original}
            if pending != set(map(str, review.get("pending_paths", []))):
                review["pending_paths"] = sorted(pending)
                self._atomic_json(self.gps_review_path, review)
        self.journal_path.unlink(missing_ok=True)

    def _recover_interrupted_discard(self, journal: dict[str, object]) -> None:
        """Completa un descarte o devuelve el original si no llegó a moverse."""
        original = str(journal["original"])
        source = Path(original)
        backup = Path(str(journal["backup"]))
        if source.is_file() and not backup.exists():
            self.journal_path.unlink(missing_ok=True)
            return
        if source.exists() or not backup.is_file():
            raise ValueError(
                "El descarte interrumpido necesita revisión manual: original y "
                "resguardo no tienen una situación inequívoca."
            )
        if original in self.completed or original in self.record_map:
            raise ValueError("El descarte interrumpido aún figura registrado en el Excel.")
        skipped = self.skipped - {original}
        last_path = "" if self.last_path == original else self.last_path
        self._write_state(skipped=skipped, last_path=last_path)
        self.skipped, self.last_path = skipped, last_path
        # La recuperación ocurre antes de cargar gps_revision.json en memoria.
        if self.gps_review_path.exists():
            review = json.loads(self.gps_review_path.read_text(encoding="utf-8"))
            pending = set(map(str, review.get("pending_paths", [])))
            if original in pending:
                review["pending_paths"] = sorted(pending - {original})
                self._atomic_json(self.gps_review_path, review)
        self.journal_path.unlink(missing_ok=True)

    def _recover_interrupted_correction(self, journal: dict[str, object]) -> None:
        original = str(journal["original"])
        old_filename = str(journal["old_filename"])
        new_filename = str(journal["new_filename"])
        new_row = dict(journal["new_row"])
        history = dict(journal["history"])
        old_destination = self.photos / old_filename
        new_destination = self.photos / new_filename
        # Comparar toda la fila evita dar por aplicada una corrección de
        # observaciones cuando el proceso se interrumpió antes de escribir Excel.
        def matches_written_row(row: dict[str, str]) -> bool:
            for key in new_row:
                expected = str(new_row.get(key, ""))
                if key not in {"Foto_en_pozo_y_nivel", "Foto_en_unidad_y_nivel"}:
                    expected = excel_safe(expected)
                if str(row.get(key, "")) != expected:
                    return False
            return True

        excel_has_new = any(matches_written_row(row) for row in self.rows)

        if excel_has_new and new_destination.is_file():
            record_map = dict(self.record_map)
            record_map[original] = new_filename
            corrections = list(self.corrections)
            if history not in corrections:
                corrections.append(history)
            self._write_state(
                last_path=original, record_map=record_map, corrections=corrections
            )
            self.record_map, self.corrections, self.last_path = record_map, corrections, original
            self.journal_path.unlink(missing_ok=True)
            return

        # Si Excel no confirmó el cambio, restaura el nombre anterior de la copia.
        if new_filename != old_filename and new_destination.exists() and not old_destination.exists():
            os.replace(new_destination, old_destination)
        self.journal_path.unlink(missing_ok=True)

    def _migrate_previous(self) -> None:
        """Importa una vez la salida anterior, sin modificarla."""
        previous_xlsx = self.output / "Catalogo_fotografico_pozos.xlsx"
        previous_csv = self.output / "Catalogo_fotografico_pozos.csv"
        old_rows: list[dict[str, str]] = []
        if previous_xlsx.exists():
            old = load_workbook(previous_xlsx, read_only=True, data_only=True).active
            if [cell.value for cell in old[1]] != OLD_FIELDS:
                raise ValueError("El Excel anterior tiene columnas diferentes; revisa la migración.")
            old_rows = [
                dict(zip(OLD_FIELDS, [str(value or "") for value in values]))
                for values in old.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in values)
            ]
        elif previous_csv.exists():
            with previous_csv.open("r", encoding="utf-8-sig", newline="") as file:
                reader = csv.DictReader(file)
                if reader.fieldnames != OLD_FIELDS:
                    raise ValueError("El CSV anterior tiene columnas diferentes; revisa la migración.")
                old_rows = list(reader)
        if not old_rows:
            return
        if not self.legacy_operation:
            raise ValueError(
                "La salida anterior no incluye Operación. Escríbela antes de migrar."
            )
        operation = self.normalize_operation(self.legacy_operation)

        for row in old_rows:
            previous_copy = self.output / row["Ruta_relativa"]
            if not previous_copy.is_file():
                raise FileNotFoundError(f"Falta la copia previa: {row['Nombre_nuevo']}")
            destination = self.photos / row["Nombre_nuevo"]
            if not destination.exists():
                shutil.copy2(previous_copy, destination)

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(FIELDS)
        for row in old_rows:
            unit_type = self.infer_legacy_unit_type(row.get("Pozo", ""))
            unit = self.normalize_unit(unit_type, row.get("Pozo", ""))
            converted = {
                "Sitio": row.get("Sitio", ""),
                "Operacion": operation,
                "Tipo_unidad": unit_type,
                "Unidad": unit,
                "Tipo_subunidad": "",
                "Subunidad": "",
                "Nivel": row.get("Nivel", ""),
                "Foto_en_unidad_y_nivel": row.get("Foto_en_pozo_y_nivel", ""),
                "Nombre_archivo": row.get("Nombre_nuevo", ""),
                "Nombre_original": row.get("Nombre_original", ""),
                "Fecha_EXIF": row.get("Fecha_EXIF", ""),
                "Vista_o_detalle": row.get("Vista_o_detalle", ""),
                "Observaciones": row.get("Observaciones", ""),
                "Fecha_catalogacion": row.get("Fecha_catalogacion", ""),
            }
            sheet.append([converted[key] for key in FIELDS])
        workbook.save(self.xlsx_path)
        completed = {row["Ruta_original"] for row in old_rows}
        record_map = {row["Ruta_original"]: row["Nombre_nuevo"] for row in old_rows}
        self._atomic_json(self.state_path, {
            "version": STATE_VERSION,
            "processed_paths": sorted(completed),
            "skipped_paths": [],
            "last_path": "",
            "record_map": record_map,
            "corrections": [],
        })

    def _append_row(self, row: dict[str, str]) -> None:
        self.ws.append([
            int(row[key]) if key == "Foto_en_unidad_y_nivel" else excel_safe(row[key])
            for key in FIELDS
        ])
        line = self.ws.max_row
        self.ws.cell(line, FIELDS.index("Foto_en_unidad_y_nivel") + 1).alignment = Alignment(
            horizontal="right"
        )

    def _write_workbook(self) -> None:
        self._assert_state_current()
        last_column = get_column_letter(len(FIELDS))
        self.ws.auto_filter.ref = f"A1:{last_column}{self.ws.max_row}"
        fd, temp_name = tempfile.mkstemp(prefix="catalogo_", suffix=".xlsx", dir=self.delivery)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            self.wb.save(temp_path)
            os.replace(temp_path, self.xlsx_path)
        finally:
            temp_path.unlink(missing_ok=True)

    @staticmethod
    def validate_identifier(value: str, label: str) -> str:
        clean = value.strip()
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,30}", clean):
            raise ValueError(f"{label} solo admite letras, números, guion y guion bajo (máximo 30).")
        return clean

    @classmethod
    def normalize_operation(cls, value: str) -> str:
        clean = value.strip()
        if clean.upper().startswith("OP") and len(clean) > 2:
            clean = clean[2:]
        return cls.validate_identifier(clean, "Operación")

    @classmethod
    def normalize_unit(cls, unit_type: str, value: str) -> str:
        if unit_type not in UNIT_TYPES:
            raise ValueError("Selecciona Pozo, Pozo auxiliar o Trinchera.")
        clean = cls.validate_identifier(value, "Código de unidad")
        upper = clean.upper()
        if unit_type == "Pozo auxiliar":
            if upper.startswith("PA") and len(clean) > 2:
                clean = clean[2:]
            elif upper.startswith("A") and len(clean) > 1:
                clean = clean[1:]
            clean = cls.validate_identifier(clean, "Código de pozo auxiliar")
            return f"A{clean}"
        if unit_type == "Trinchera":
            if upper.startswith("TR") and len(clean) > 2:
                clean = clean[2:]
            elif upper.startswith("T") and len(clean) > 1:
                clean = clean[1:]
            clean = cls.validate_identifier(clean, "Código de trinchera")
            return f"T{clean}"
        if upper.startswith("P") and len(clean) > 1:
            clean = clean[1:]
        return cls.validate_identifier(clean, "Código de pozo")

    @staticmethod
    def unit_form_value(unit_type: str, value: str) -> str:
        """Muestra el número de trinchera sin el prefijo guardado en Excel."""
        if unit_type == "Trinchera":
            clean = value.strip()
            upper = clean.upper()
            if upper.startswith("TR") and len(clean) > 2:
                return clean[2:]
            if upper.startswith("T") and len(clean) > 1:
                return clean[1:]
        return value

    @classmethod
    def normalize_subunit(cls, subunit_type: str, value: str) -> tuple[str, str]:
        """Valida una subunidad opcional; ambos campos se usan juntos o quedan vacíos."""
        clean_type = subunit_type.strip()
        clean_value = value.strip()
        if not clean_type and not clean_value:
            return "", ""
        if not clean_type or not clean_value:
            raise ValueError(
                "Para usar una subunidad completa tanto el tipo como su código o número."
            )
        if clean_type not in SUBUNIT_TYPES or not clean_type:
            raise ValueError("Selecciona un tipo de subunidad de la lista.")
        return clean_type, cls.validate_identifier(clean_value, "Código de subunidad")

    @staticmethod
    def infer_legacy_unit_type(value: str) -> str:
        upper = value.strip().upper()
        if upper.startswith("A") or upper.startswith("PA"):
            return "Pozo auxiliar"
        if upper.startswith("T"):
            return "Trinchera"
        return "Pozo"

    @staticmethod
    def unit_filename_token(unit_type: str, unit: str) -> str:
        return unit if unit_type == "Trinchera" else f"P{unit}"

    @staticmethod
    def subunit_filename_token(subunit_type: str, subunit: str) -> str:
        if not subunit_type:
            return ""
        return f"{SUBUNIT_FILENAME_CODES[subunit_type]}{subunit}"

    def _migrate_operation_and_units(self, legacy_operation: str | None) -> None:
        if not legacy_operation:
            raise ValueError(
                "Este catálogo fue creado antes de incluir Operación. Escribe la "
                "operación de las filas existentes y vuelve a cargarlo."
            )
        operation = self.normalize_operation(legacy_operation)
        migrated: list[dict[str, str]] = []
        for old in self.rows:
            unit_type = self.infer_legacy_unit_type(old["Pozo"])
            unit = self.normalize_unit(unit_type, old["Pozo"])
            migrated.append({
                "Sitio": old["Sitio"],
                "Operacion": operation,
                "Tipo_unidad": unit_type,
                "Unidad": unit,
                "Tipo_subunidad": "",
                "Subunidad": "",
                "Nivel": old["Nivel"],
                "Foto_en_unidad_y_nivel": old["Foto_en_pozo_y_nivel"],
                "Nombre_archivo": old["Nombre_archivo"],
                "Nombre_original": old["Nombre_original"],
                "Fecha_EXIF": old["Fecha_EXIF"],
                "Vista_o_detalle": old["Vista_o_detalle"],
                "Observaciones": old["Observaciones"],
                "Fecha_catalogacion": old["Fecha_catalogacion"],
            })

        backup_xlsx = self.output / "respaldo_catalogo_antes_v2_2.xlsx"
        backup_state = self.output / "respaldo_estado_antes_v2_2.json"
        if not backup_xlsx.exists():
            shutil.copy2(self.xlsx_path, backup_xlsx)
        if self.state_path.exists() and not backup_state.exists():
            shutil.copy2(self.state_path, backup_state)

        old_wb, old_ws, old_rows = self.wb, self.ws, self.rows
        try:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Fotografías"
            self.ws.append(FIELDS)
            self.ws.freeze_panes = "H2"
            self.ws.sheet_view.showGridLines = False
            self.ws.row_dimensions[1].height = 31
            for cell in self.ws[1]:
                cell.fill = PatternFill("solid", fgColor="2F5D50")
                cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", horizontal="center")
            widths = [14, 12, 17, 15, 18, 16, 12, 24, 48, 35, 23, 34, 52, 22]
            for index, width in enumerate(widths, 1):
                self.ws.column_dimensions[get_column_letter(index)].width = width
            self.rows = []
            for row in migrated:
                self._append_row(row)
            self._write_workbook()
            self.rows = migrated
            self.needs_schema_migration = False
        except Exception:
            self.wb, self.ws, self.rows = old_wb, old_ws, old_rows
            raise
        try:
            self._write_state()
        except OSError:
            pass

    def _migrate_optional_subunits(self) -> None:
        """Añade columnas opcionales sin cambiar filas ni nombres ya registrados."""
        migrated = [
            {key: str(old.get(key, "")) for key in FIELDS}
            for old in self.rows
        ]
        backup_xlsx = self.output / "respaldo_catalogo_antes_v2_3.xlsx"
        backup_state = self.output / "respaldo_estado_antes_v2_3.json"
        if not backup_xlsx.exists():
            shutil.copy2(self.xlsx_path, backup_xlsx)
        if self.state_path.exists() and not backup_state.exists():
            shutil.copy2(self.state_path, backup_state)

        old_wb, old_ws, old_rows = self.wb, self.ws, self.rows
        try:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Fotografías"
            self.ws.append(FIELDS)
            self.ws.freeze_panes = "H2"
            self.ws.sheet_view.showGridLines = False
            self.ws.row_dimensions[1].height = 31
            for cell in self.ws[1]:
                cell.fill = PatternFill("solid", fgColor="2F5D50")
                cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", horizontal="center")
            widths = [14, 12, 17, 15, 18, 16, 12, 24, 48, 35, 23, 34, 52, 22]
            for index, width in enumerate(widths, 1):
                self.ws.column_dimensions[get_column_letter(index)].width = width
            self.rows = []
            for row in migrated:
                self._append_row(row)
            self._write_workbook()
            self.rows = migrated
            self.needs_subunit_migration = False
        except Exception:
            self.wb, self.ws, self.rows = old_wb, old_ws, old_rows
            raise
        try:
            self._write_state()
        except OSError:
            pass

    def remaining(self) -> list[Path]:
        return [path for path in self.files
                if normalized(path) not in self.completed and normalized(path) not in self.skipped]

    def omitted(self) -> list[Path]:
        return [path for path in self.files if normalized(path) in self.skipped]

    def status_of(self, path: Path) -> str:
        value = normalized(path)
        if value in self.completed:
            return "registered"
        if value in self.skipped:
            return "skipped"
        return "pending"

    def filtered_files(self, condition: str) -> list[Path]:
        """Filtra para navegar; nunca modifica el estado ni el Excel."""
        if condition not in PHOTO_FILTERS:
            raise ValueError(f"Condición de filtro desconocida: {condition}")
        if condition == "Todas":
            return list(self.files)
        if condition == "Por catalogar":
            return [path for path in self.files if self.status_of(path) != "registered"]
        if condition == "Pendientes":
            return self.remaining()
        if condition == "Omitidas":
            return self.omitted()
        if condition == "Registradas":
            return [path for path in self.files if self.status_of(path) == "registered"]
        return [path for path in self.files if self.display_status(path) == "gps_review"]

    def _rebuild_record_map(self) -> None:
        """Completa asociaciones inequívocas de estados creados por versiones anteriores."""
        filenames = {row["Nombre_archivo"] for row in self.rows}
        self.record_map = {
            path: filename for path, filename in self.record_map.items()
            if path in self.completed and filename in filenames
        }
        rows_by_original: dict[str, list[dict[str, str]]] = {}
        paths_by_name: dict[str, list[str]] = {}
        for row in self.rows:
            rows_by_original.setdefault(row["Nombre_original"], []).append(row)
        for path in self.completed:
            paths_by_name.setdefault(Path(path).name, []).append(path)
        for name, paths in paths_by_name.items():
            candidates = rows_by_original.get(name, [])
            if len(paths) == 1 and len(candidates) == 1:
                self.record_map.setdefault(paths[0], candidates[0]["Nombre_archivo"])

    def row_for_path(self, path: Path) -> tuple[int, dict[str, str]] | None:
        original = normalized(path)
        filename = self.record_map.get(original)
        if filename:
            for index, row in enumerate(self.rows):
                if row["Nombre_archivo"] == filename:
                    return index, row
        matches = [
            (index, row) for index, row in enumerate(self.rows)
            if row["Nombre_original"] == path.name
        ]
        completed_same_name = [value for value in self.completed if Path(value).name == path.name]
        if len(matches) == 1 and len(completed_same_name) == 1:
            self.record_map[original] = matches[0][1]["Nombre_archivo"]
            return matches[0]
        return None

    def next_number(
        self, operation: str, unit_type: str, unit: str, level: str,
        subunit_type: str = "", subunit: str = "",
    ) -> int:
        numbers = [int(row["Foto_en_unidad_y_nivel"]) for row in self.rows
                   if row["Operacion"] == operation
                   and row["Tipo_unidad"] == unit_type
                   and row["Unidad"] == unit
                   and row.get("Tipo_subunidad", "") == subunit_type
                   and row.get("Subunidad", "") == subunit
                   and row["Nivel"] == level
                   and str(row["Foto_en_unidad_y_nivel"]).isdigit()]
        return max(numbers, default=0) + 1

    def proposed_name(
        self, path: Path, operation: str, unit_type: str, unit: str, level: str,
        subunit_type: str = "", subunit: str = "",
    ) -> str:
        operation = self.normalize_operation(operation)
        unit = self.normalize_unit(unit_type, unit)
        subunit_type, subunit = self.normalize_subunit(subunit_type, subunit)
        level = self.validate_identifier(level, "Nivel")
        safe_site = self.validate_identifier(self.site, "Código del sitio")
        number = self.next_number(
            operation, unit_type, unit, level, subunit_type, subunit
        )
        unit_token = self.unit_filename_token(unit_type, unit)
        subunit_token = self.subunit_filename_token(subunit_type, subunit)
        context = f"{unit_token}_{subunit_token}" if subunit_token else unit_token
        return (
            f"{safe_site}_OP{operation}_{context}_{level}_"
            f"F{number:03d}{path.suffix.lower()}"
        )

    def proposed_correction_name(
        self, path: Path, operation: str, unit_type: str, unit: str, level: str,
        subunit_type: str = "", subunit: str = "",
    ) -> str:
        operation = self.normalize_operation(operation)
        unit = self.normalize_unit(unit_type, unit)
        subunit_type, subunit = self.normalize_subunit(subunit_type, subunit)
        level = self.validate_identifier(level, "Nivel")
        located = self.row_for_path(path)
        if located is None:
            raise ValueError(
                "No se pudo asociar esta fotografía con una única fila del Excel. "
                "Puede ocurrir con catálogos antiguos que contienen nombres originales repetidos."
            )
        _, old_row = located
        same_group = (
            old_row["Operacion"] == operation
            and old_row["Tipo_unidad"] == unit_type
            and old_row["Unidad"] == unit
            and old_row.get("Tipo_subunidad", "") == subunit_type
            and old_row.get("Subunidad", "") == subunit
            and old_row["Nivel"] == level
        )
        if same_group:
            return old_row["Nombre_archivo"]
        number = self.next_number(
            operation, unit_type, unit, level, subunit_type, subunit
        )
        safe_site = self.validate_identifier(self.site, "Código del sitio")
        unit_token = self.unit_filename_token(unit_type, unit)
        subunit_token = self.subunit_filename_token(subunit_type, subunit)
        context = f"{unit_token}_{subunit_token}" if subunit_token else unit_token
        return (
            f"{safe_site}_OP{operation}_{context}_{level}_"
            f"F{number:03d}{path.suffix.lower()}"
        )

    def _replace_worksheet_row(self, row_index: int, row: dict[str, str]) -> None:
        excel_row = row_index + 2
        for column, key in enumerate(FIELDS, 1):
            value: object = int(row[key]) if key == "Foto_en_unidad_y_nivel" else excel_safe(row[key])
            self.ws.cell(excel_row, column, value)
        self.ws.cell(excel_row, FIELDS.index("Foto_en_unidad_y_nivel") + 1).alignment = Alignment(
            horizontal="right"
        )

    def set_skipped(self, path: Path, skipped: bool) -> None:
        original = normalized(path)
        if original in self.completed:
            raise ValueError("La fotografía ya está registrada y no se puede omitir.")
        new_skipped = set(self.skipped)
        if skipped:
            new_skipped.add(original)
        else:
            new_skipped.discard(original)
        self._write_state(skipped=new_skipped, last_path=original)
        self.skipped = new_skipped
        self.last_path = original

    def withdraw_for_review(self, path: Path) -> str:
        """Quita la fila y la copia; conserva el original como omitido por revisar."""
        original = normalized(path)
        if original not in self.completed:
            raise ValueError("Esta fotografía no está registrada.")
        located = self.row_for_path(path)
        if located is None:
            raise ValueError("No se pudo asociar la fotografía con una única fila.")
        row_index, row = located
        filename = row["Nombre_archivo"]
        if Path(filename).name != filename:
            raise ValueError("El nombre de la copia no es seguro.")
        destination = self.photos / filename
        if not destination.is_file():
            raise FileNotFoundError(f"No existe la copia {filename}.")
        backup_dir = self.output / "Respaldos_retiros"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{filename}"
        if backup.exists():
            raise FileExistsError("Ya existe el resguardo de esta copia.")
        self._write_journal({
            "version": 1, "operation": "withdraw", "original": original,
            "filename": filename, "backup": str(backup),
            "started_at": datetime.now().isoformat(timespec="seconds"),
        })
        workbook_written = False
        try:
            os.replace(destination, backup)
            self.ws.delete_rows(row_index + 2)
            self._write_workbook()
            workbook_written = True
            completed = self.completed - {original}
            skipped = self.skipped | {original}
            record_map = dict(self.record_map)
            record_map.pop(original, None)
            self.rows.pop(row_index)
            try:
                self._write_state(
                    completed=completed, skipped=skipped, last_path=original,
                    record_map=record_map,
                )
            except Exception as exc:
                self.completed, self.skipped, self.last_path = completed, skipped, original
                self.record_map = record_map
                raise WithdrawalCommittedWarning(
                    "La copia salió de la entrega y la fila se quitó, pero el estado se "
                    "completará al reabrir. No repitas la operación."
                ) from exc
            self.completed, self.skipped, self.last_path = completed, skipped, original
            self.record_map = record_map
            if original in self.gps_review:
                try:
                    pending = self.gps_review - {original}
                    self._write_gps_review(pending, self.gps_manual_history)
                    self.gps_review = pending
                except Exception as exc:
                    raise WithdrawalCommittedWarning(
                        "El registro se retiró, pero la lista GPS se terminará de "
                        "actualizar al reabrir. No repitas la operación."
                    ) from exc
            self.journal_path.unlink(missing_ok=True)
            return filename
        except WithdrawalCommittedWarning:
            raise
        except Exception:
            if not workbook_written:
                self.wb = load_workbook(self.xlsx_path)
                self.ws = self.wb.active
                if backup.exists() and not destination.exists():
                    os.replace(backup, destination)
                self.journal_path.unlink(missing_ok=True)
            raise

    def discard_unregistered(self, path: Path) -> Path:
        """Saca un original pendiente/omitido del trabajo y lo deja recuperable."""
        original = normalized(path)
        if original in self.completed or original in self.record_map:
            raise ValueError("Retira primero esta foto registrada del catálogo.")
        if path not in self.files or not path.is_file():
            raise FileNotFoundError("La fotografía original no está disponible.")
        if self.row_for_path(path) is not None:
            raise ValueError("Esta fotografía todavía tiene una fila de Excel.")
        backup_dir = self.output / "Descartadas_catalogador" / (
            datetime.now().strftime("%Y%m%d_%H%M%S_%f") + "_" + path.stem
        )
        backup_dir.mkdir(parents=True, exist_ok=False)
        backup = backup_dir / path.name
        self._atomic_json(backup_dir / "origen.json", {
            "original": original,
            "status_before": "skipped" if original in self.skipped else "pending",
            "discarded_at": datetime.now().isoformat(timespec="seconds"),
        })
        self._write_journal({
            "version": 1, "operation": "discard_unregistered",
            "original": original, "backup": str(backup),
        })
        state_written = False
        try:
            shutil.move(str(path), str(backup))
            if path.exists() or not backup.is_file():
                raise OSError("No se pudo comprobar el resguardo del original.")
            skipped = self.skipped - {original}
            last_path = "" if self.last_path == original else self.last_path
            self._write_state(skipped=skipped, last_path=last_path)
            state_written = True
            self.skipped, self.last_path = skipped, last_path
            self.files = [item for item in self.files if normalized(item) != original]
            if original in self.gps_review:
                pending = self.gps_review - {original}
                self._write_gps_review(pending, self.gps_manual_history)
                self.gps_review = pending
            self.journal_path.unlink(missing_ok=True)
            return backup
        except Exception as exc:
            if not state_written and self._state_digest == self._disk_state_digest():
                if backup.is_file() and not path.exists():
                    shutil.move(str(backup), str(path))
                self.journal_path.unlink(missing_ok=True)
                raise
            self.skipped.discard(original)
            self.files = [item for item in self.files if normalized(item) != original]
            raise DiscardCommittedWarning(
                "El original quedó resguardado, pero hay una actualización pendiente. "
                "Cierra y reabre el catálogo antes de continuar."
            ) from exc

    def _disk_state_digest(self) -> bytes | None:
        return (hashlib.sha256(self.state_path.read_bytes()).digest()
                if self.state_path.exists() else None)

    def save(
        self, path: Path, operation: str, unit_type: str, unit: str,
        level: str, view: str, notes: str,
        subunit_type: str = "", subunit: str = "",
        manual_gps: tuple[str, str, str, str] | None = None,
    ) -> dict[str, str]:
        operation = self.normalize_operation(operation)
        unit = self.normalize_unit(unit_type, unit)
        subunit_type, subunit = self.normalize_subunit(subunit_type, subunit)
        level = self.validate_identifier(level, "Nivel")
        original = normalized(path)
        if original in self.completed:
            raise ValueError("Esta fotografía ya está registrada.")

        number = self.next_number(
            operation, unit_type, unit, level, subunit_type, subunit
        )
        filename = self.proposed_name(
            path, operation, unit_type, unit, level, subunit_type, subunit
        )
        destination = self.photos / filename
        if destination.exists():
            raise FileExistsError(f"Ya existe {filename}. Revisa la carpeta antes de continuar.")
        if manual_gps is not None:
            if destination.suffix.lower() not in {".jpg", ".jpeg"}:
                raise ValueError("El GPS manual sin recomprimir solo está disponible para JPEG.")
            if read_gps(path)[0] is not None:
                raise ValueError("El original ya tiene GPS válido; no se sustituyó.")

        exif_date = ""
        try:
            with Image.open(path) as image:
                exif = image.getexif()
                exif_date = str(exif.get(36867) or exif.get(306) or "")
        except Exception:
            pass

        row = dict(zip(FIELDS, [
            self.site, operation, unit_type, unit, subunit_type, subunit, level,
            str(number), filename,
            path.name, exif_date, view.strip(), notes.strip(),
            datetime.now().isoformat(timespec="seconds"),
        ]))
        self._write_journal({
            "version": 1, "operation": "save", "original": original, "filename": filename,
            "started_at": datetime.now().isoformat(timespec="seconds"),
        })
        copied = False
        row_appended = False
        workbook_written = False
        try:
            shutil.copy2(path, destination)
            copied = True
            if manual_gps is not None:
                self._replace_copy_bytes(
                    destination, add_manual_gps(destination.read_bytes(), *manual_gps)
                )
                if read_gps(destination)[0] is None:
                    raise RuntimeError("No se pudo comprobar el GPS manual de la copia.")
            self._append_row(row)
            row_appended = True
            self._write_workbook()
            workbook_written = True

            completed = self.completed | {original}
            skipped = self.skipped - {original}
            record_map = dict(self.record_map)
            record_map[original] = filename
            try:
                self._write_state(
                    completed=completed, skipped=skipped, last_path=original,
                    record_map=record_map,
                )
            except Exception as exc:
                self.rows.append(row)
                self.completed, self.skipped, self.last_path = completed, skipped, original
                self.record_map = record_map
                raise SaveCommittedWarning(
                    "La copia y la fila de Excel se guardaron. No se pudo actualizar el estado; "
                    "no repitas la foto. Al reabrir, el programa intentará recuperarlo."
                ) from exc

            self.rows.append(row)
            self.completed, self.skipped, self.last_path = completed, skipped, original
            self.record_map = record_map
            self.journal_path.unlink(missing_ok=True)
            try:
                pending_gps = set(self.gps_review)
                history_gps = list(self.gps_manual_history)
                if manual_gps is not None:
                    pending_gps.discard(original)
                    history_gps.append({
                        "original": original, "filename": filename,
                        "latitude": manual_gps[0].strip(),
                        "longitude": manual_gps[1].strip(),
                        "altitude_m": manual_gps[2].strip(),
                        "accuracy_m": manual_gps[3].strip(),
                        "source": "transcripción manual de coordenadas de cámara",
                        "recorded_at": datetime.now().isoformat(timespec="seconds"),
                    })
                elif read_gps(destination)[0] is None:
                    pending_gps.add(original)
                if pending_gps != self.gps_review or history_gps != self.gps_manual_history:
                    self._write_gps_review(pending_gps, history_gps)
                    self.gps_review, self.gps_manual_history = pending_gps, history_gps
            except Exception as exc:
                raise SaveCommittedWarning(
                    "La foto, el Excel y el estado sí se guardaron, pero no se pudo actualizar "
                    "la lista GPS por revisar. No repitas el registro."
                ) from exc
            return row
        except SaveCommittedWarning:
            raise
        except Exception:
            if row_appended and not workbook_written and self.ws.max_row > len(self.rows) + 1:
                self.ws.delete_rows(self.ws.max_row)
            if copied and not workbook_written:
                destination.unlink(missing_ok=True)
            if not workbook_written:
                self.journal_path.unlink(missing_ok=True)
            raise

    def correct(
        self, path: Path, operation: str, unit_type: str, unit: str,
        level: str, view: str, notes: str,
        subunit_type: str = "", subunit: str = "",
    ) -> dict[str, str]:
        """Corrige una fila y el nombre de su copia sin modificar el original."""
        operation = self.normalize_operation(operation)
        unit = self.normalize_unit(unit_type, unit)
        subunit_type, subunit = self.normalize_subunit(subunit_type, subunit)
        level = self.validate_identifier(level, "Nivel")
        original = normalized(path)
        if original not in self.completed:
            raise ValueError("Solo se pueden corregir fotografías ya registradas.")
        located = self.row_for_path(path)
        if located is None:
            raise ValueError(
                "No se pudo identificar de forma inequívoca la fila de esta fotografía."
            )
        row_index, old_row = located
        old_row = dict(old_row)
        old_filename = old_row["Nombre_archivo"]
        new_filename = self.proposed_correction_name(
            path, operation, unit_type, unit, level, subunit_type, subunit
        )
        old_destination = self.photos / old_filename
        new_destination = self.photos / new_filename
        if not old_destination.is_file():
            raise FileNotFoundError(f"Falta la copia registrada: {old_filename}")
        if new_filename != old_filename and new_destination.exists():
            raise FileExistsError(f"Ya existe {new_filename}; no se aplicó la corrección.")

        if (
            old_row["Operacion"] == operation
            and old_row["Tipo_unidad"] == unit_type
            and old_row["Unidad"] == unit
            and old_row.get("Tipo_subunidad", "") == subunit_type
            and old_row.get("Subunidad", "") == subunit
            and old_row["Nivel"] == level
        ):
            number = old_row["Foto_en_unidad_y_nivel"]
        else:
            number = str(self.next_number(
                operation, unit_type, unit, level, subunit_type, subunit
            ))
        new_row = dict(old_row)
        new_row.update({
            "Operacion": operation,
            "Tipo_unidad": unit_type,
            "Unidad": unit,
            "Tipo_subunidad": subunit_type,
            "Subunidad": subunit,
            "Nivel": level,
            "Foto_en_unidad_y_nivel": str(number),
            "Nombre_archivo": new_filename,
            "Vista_o_detalle": view.strip(),
            "Observaciones": notes.strip(),
        })
        history: dict[str, object] = {
            "original_path": original,
            "corrected_at": datetime.now().isoformat(timespec="seconds"),
            "before": old_row,
            "after": new_row,
        }
        self._write_journal({
            "version": 1,
            "operation": "correction",
            "original": original,
            "old_filename": old_filename,
            "new_filename": new_filename,
            "old_row": old_row,
            "new_row": new_row,
            "history": history,
        })

        renamed = False
        workbook_written = False
        try:
            if new_filename != old_filename:
                os.replace(old_destination, new_destination)
                renamed = True
            self._replace_worksheet_row(row_index, new_row)
            self._write_workbook()
            workbook_written = True
            record_map = dict(self.record_map)
            record_map[original] = new_filename
            corrections = list(self.corrections) + [history]
            try:
                self._write_state(
                    last_path=original, record_map=record_map, corrections=corrections
                )
            except Exception as exc:
                self.rows[row_index] = new_row
                self.record_map, self.corrections, self.last_path = record_map, corrections, original
                raise CorrectionCommittedWarning(
                    "La corrección quedó aplicada en la copia y el Excel. El estado se "
                    "completará automáticamente al reabrir."
                ) from exc
            self.rows[row_index] = new_row
            self.record_map, self.corrections, self.last_path = record_map, corrections, original
            self.journal_path.unlink(missing_ok=True)
            return new_row
        except CorrectionCommittedWarning:
            raise
        except Exception:
            if not workbook_written:
                self._replace_worksheet_row(row_index, old_row)
                if renamed and new_destination.exists() and not old_destination.exists():
                    os.replace(new_destination, old_destination)
                self.journal_path.unlink(missing_ok=True)
            raise
