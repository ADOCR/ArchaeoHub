"""Dominio local de inventario. Una solicitud por proceso; SQLite independiente del Hub.

Las filas normalizadas importadas se conservan sin reclasificar. Las exportaciones no sustituyen archivos.
"""
from pathlib import Path
import csv
import hashlib
import io
import json
import math
import os
import sqlite3
import sys
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timezone

CATEGORIES = {'q_arc':'Arcilla','q_arc_coc':'Arcilla cocida','q_cer':'Cerámica',
              'q_ind':'Indeterminado','q_lit':'Lítica','q_const':'Material constructivo',
              'q_muestra':'Muestra','q_vidrio':'Vidrio'}
COLLECTIONS = ['fragmentaria','artefactos','humanos','fauna','botanica','carbon','pendiente']
SCHEMA = 1
MAX_ROWS = 100000

def stamp(): return datetime.now(timezone.utc).isoformat()
def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()
def dumps(value): return json.dumps(value,ensure_ascii=False,allow_nan=False,default=str)
def text(value): return '' if value is None else str(value).strip()
def yes(value): return text(value).lower() in ('sí','si','yes','true','1') or value is True
def number(value, label='Cantidad'):
    if value is None or text(value)=='':return None
    if isinstance(value,bool):raise ValueError(f'{label}: escribe un número, no un booleano.')
    try:n=float(value)
    except (TypeError,ValueError):raise ValueError(f'{label}: valor no numérico: {value}')
    if not math.isfinite(n) or n<0 or n>1e12:raise ValueError(f'{label}: valor fuera de rango.')
    return int(n) if n.is_integer() else n
def count(value,label='Cantidad'):
    n=number(value,label)
    if n is not None and not isinstance(n,int):raise ValueError(f'{label}: el conteo debe ser entero.')
    return n
def cancelled(request):
    if request.get('cancelFile') and Path(request['cancelFile']).exists():raise ValueError('Operación cancelada; no se aplicaron cambios incompletos.')
def emit(value):
    sys.stdout.buffer.write((dumps(value)+'\n').encode('utf-8'));sys.stdout.buffer.flush()
def progress(value,message):emit({'type':'progress','progress':value,'message':message})

@contextmanager
def database(path):
    p=Path(path)
    if not p.is_file() or p.is_symlink():raise ValueError('No se encuentra el inventario seleccionado.')
    con=sqlite3.connect(p,timeout=5)
    try:
        con.execute('PRAGMA foreign_keys=ON')
        if con.execute('PRAGMA application_id').fetchone()[0]!=0x4148494E or con.execute('PRAGMA user_version').fetchone()[0]!=SCHEMA:
            raise ValueError('El archivo no es un inventario compatible de ArchaeoHub.')
        yield con
    finally:con.close()

def create(request):
    p=Path(request['database']);metadata=request.get('metadata',{})
    if not text(metadata.get('project')) or not text(metadata.get('site')):raise ValueError('Indica proyecto y código del sitio.')
    # Reserva exclusiva: nunca inicialices una base arqueológica existente.
    with p.open('xb'):pass
    try:
        con=sqlite3.connect(p)
        try:
            con.executescript('''PRAGMA application_id=1095256398; PRAGMA user_version=1;
            PRAGMA foreign_keys=ON;
            CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
            CREATE TABLE sources(id INTEGER PRIMARY KEY,path TEXT NOT NULL,sha256 TEXT UNIQUE NOT NULL,imported_at TEXT NOT NULL);
            CREATE TABLE contexts(id INTEGER PRIMARY KEY,signature TEXT UNIQUE NOT NULL,data TEXT NOT NULL);
            CREATE TABLE containers(id INTEGER PRIMARY KEY,parent_id INTEGER REFERENCES containers(id),kind TEXT NOT NULL,code TEXT NOT NULL,signature TEXT UNIQUE NOT NULL);
            CREATE TABLE records(id TEXT PRIMARY KEY,context_id INTEGER NOT NULL REFERENCES contexts(id),container_id INTEGER REFERENCES containers(id),source_id INTEGER REFERENCES sources(id),collection TEXT NOT NULL,data TEXT NOT NULL,raw_data TEXT NOT NULL,revision INTEGER NOT NULL);
            CREATE TABLE materials(record_id TEXT REFERENCES records(id),field TEXT NOT NULL,quantity REAL,PRIMARY KEY(record_id,field));
            CREATE TABLE audit(id INTEGER PRIMARY KEY,record_id TEXT NOT NULL,action TEXT NOT NULL,at TEXT NOT NULL,before_data TEXT,after_data TEXT NOT NULL);
            ''')
            con.executemany('INSERT INTO metadata VALUES(?,?)',[(k,dumps(v)) for k,v in metadata.items()])
            con.commit()
        finally:con.close()
    except Exception:
        p.unlink(missing_ok=True);raise
    return snapshot(request)

CONTEXT_FIELDS=('operacion','pozo','contexto','trinchera','unidad','nivel_std','prof_desde_cm','prof_hasta_cm')
def put_record(con,data,raw,source_id=None,revision=1):
    rid=text(data.get('registro'))
    if not rid or len(rid)>256:raise ValueError('Cada registro requiere un identificador estable de hasta 256 caracteres.')
    for key in ('cantidad_original','cantidad_analisis','suma_desglose',*CATEGORIES):number(data.get(key),key)
    context={k:data.get(k) for k in CONTEXT_FIELDS}
    signature=dumps(context)
    con.execute('INSERT OR IGNORE INTO contexts(signature,data) VALUES(?,?)',(signature,signature))
    context_id=con.execute('SELECT id FROM contexts WHERE signature=?',(signature,)).fetchone()[0]
    parent=None;parts=[]
    for kind,field in [('box','caja'),('large_bag','bolsa_grande'),('small_bag','bolsita')]:
        code=text(data.get(field))
        if not code:continue
        parts.append([kind,code]);signature=dumps(parts)
        con.execute('INSERT OR IGNORE INTO containers(parent_id,kind,code,signature) VALUES(?,?,?,?)',(parent,kind,code,signature))
        parent=con.execute('SELECT id FROM containers WHERE signature=?',(signature,)).fetchone()[0]
    collection=data.get('collection_type') or ('artefactos' if yes(data.get('es_artefacto')) else 'pendiente' if source_id else 'fragmentaria')
    if collection not in COLLECTIONS:raise ValueError('Tipo de colección no reconocido.')
    con.execute('INSERT INTO records VALUES(?,?,?,?,?,?,?,?)',(rid,context_id,parent,source_id,collection,dumps(data),dumps(raw),revision))
    con.executemany('INSERT INTO materials VALUES(?,?,?)',[(rid,k,number(v,k)) for k,v in data.items() if k in CATEGORIES or k.startswith('n_')])
    con.execute('INSERT INTO audit(record_id,action,at,after_data) VALUES(?,?,?,?)',(rid,'import' if source_id else 'create',stamp(),dumps(data)))

def load_source(path):
    p=Path(path)
    if not p.is_file() or p.is_symlink() or p.stat().st_size>64*1024*1024:raise ValueError('Selecciona una fuente local de hasta 64 MB.')
    sha=digest(p)
    if p.suffix.lower()=='.json':
        doc=json.loads(p.read_text(encoding='utf-8-sig'))
        rows=doc.get('inventario') if isinstance(doc,dict) else doc
        headers=doc.get('inv_headers') if isinstance(doc,dict) else None
        if not isinstance(rows,list):raise ValueError('El JSON debe contener inventario normalizado; no se clasifican descripciones automáticamente.')
        if rows and isinstance(rows[0],list):
            if not headers or any(len(row)!=len(headers) for row in rows):raise ValueError('Las filas no coinciden con los encabezados declarados.')
            rows=[dict(zip(headers,row)) for row in rows]
    elif p.suffix.lower() in ('.csv','.xlsx'):
        if p.suffix.lower()=='.xlsx':
            with zipfile.ZipFile(p) as archive:
                if len(archive.infolist())>4096 or sum(x.file_size for x in archive.infolist())>128*1024*1024:
                    raise ValueError('El contenido expandido del Excel supera 128 MB o 4.096 entradas.')
            from openpyxl import load_workbook
            book=load_workbook(p,read_only=True,data_only=False,keep_links=False)
            try:
                sheets=[s for s in book if s.title.lower() in ('inventario','registros')]
                if len(sheets)!=1:raise ValueError('El Excel requiere una hoja Inventario o Registros normalizada. Conserva el original y usa su versión procesada.')
                if sheets[0].max_row>MAX_ROWS+1 or sheets[0].max_column>256:raise ValueError('La hoja supera 100.000 filas o 256 columnas.')
                values=list(sheets[0].values)
                if any(isinstance(x,str) and x.startswith('=') for row in values for x in row):raise ValueError('Hay fórmulas; exporta primero una copia con valores verificados.')
            finally:book.close()
        else:
            values=[]
            with p.open(encoding='utf-8-sig',newline='') as f:
                for row in csv.reader(f):
                    if len(values)>MAX_ROWS or len(row)>256:raise ValueError('El CSV supera 100.000 filas o 256 columnas.')
                    values.append(row)
        if not values:raise ValueError('La tabla está vacía.')
        headers=[text(x) for x in values[0]]
        if len(set(headers))!=len(headers):raise ValueError('Hay encabezados repetidos.')
        rows=[dict(zip(headers,row)) for row in values[1:] if any(x is not None and text(x) for x in row)]
    else:raise ValueError('Importa JSON, CSV o Excel XLSX normalizado.')
    if not rows or len(rows)>MAX_ROWS or any(not isinstance(row,dict) for row in rows):raise ValueError('La fuente debe contener entre 1 y 100.000 registros.')
    if any(len(row)>256 for row in rows):raise ValueError('La fuente supera 256 campos por registro.')
    if any(not {'registro','cantidad_original','cantidad_analisis','usar_analisis'}.issubset(row) for row in rows):
        raise ValueError('Faltan campos normalizados: registro, cantidad_original, cantidad_analisis y usar_analisis. No se inventarán cantidades.')
    if len({text(r['registro']) for r in rows})!=len(rows):raise ValueError('La fuente repite identificadores de registro.')
    if digest(p)!=sha:raise ValueError('La fuente cambió durante la lectura. Revisa nuevamente.')
    return rows,sha

def usable(row):return yes(row.get('usar_analisis')) and not text(row.get('duplicado_de')) and number(row.get('cantidad_analisis')) is not None
def statistics(rows):
    valid=[r for r in rows if usable(r)]
    groups={}
    for row in valid:
        key=tuple(text(row.get(k)) for k in ('operacion','pozo','nivel_std'))
        if key not in groups:groups[key]={'operacion':key[0],'pozo':key[1],'nivel_std':key[2],'cantidad':0,'registros':0,**{k:0 for k in CATEGORIES}}
        item=groups[key];item['cantidad']+=number(row['cantidad_analisis']);item['registros']+=1
        for k in CATEGORIES:item[k]+=number(row.get(k)) or 0
    return {'records':len(rows),'usable':len(valid),'pending':sum('Diferencia' in text(r.get('cantidad_estado')) for r in rows),
            'excluded':len(rows)-len(valid),'duplicates':sum(bool(text(r.get('duplicado_de'))) for r in rows),
            'quantity':sum(number(r['cantidad_analisis']) for r in valid),'categories':{k:sum(number(r.get(k)) or 0 for r in valid) for k in CATEGORIES},
            'groups':list(groups.values())}

def preview(request):
    rows,sha=load_source(request['source']);result=statistics(rows)
    with database(request['database']) as con:
        exists={x[0] for x in con.execute('SELECT id FROM records')}
        result['collisions']=[text(r['registro']) for r in rows if text(r['registro']) in exists][:30]
        result['alreadyImported']=bool(con.execute('SELECT 1 FROM sources WHERE sha256=?',(sha,)).fetchone())
    result.update({'sourceDigest':sha,'source':request['source'],'sample':rows[:5]})
    return result

def import_source(request):
    rows,sha=load_source(request['source'])
    if sha!=request.get('sourceDigest'):raise ValueError('La fuente cambió desde la vista previa. Revísala nuevamente.')
    with database(request['database']) as con,con:
        source=con.execute('INSERT INTO sources(path,sha256,imported_at) VALUES(?,?,?)',(str(Path(request['source']).resolve()),sha,stamp())).lastrowid
        for i,row in enumerate(rows):
            cancelled(request);put_record(con,row,row,source)
            if i%200==0:progress(i/len(rows),'Importando registros sin reclasificar la fuente…')
        cancelled(request)
    return snapshot(request)

def prepare_manual(data):
    row=dict(data);row['registro']=text(row.get('registro')) or 'manual:'+str(uuid.uuid4())
    if not text(row.get('operacion')):raise ValueError('Indica la operación del contexto.')
    collection=row.get('collection_type','fragmentaria');row['collection_type']=collection
    for k in ('prof_desde_cm','prof_hasta_cm'):row[k]=number(row.get(k),k)
    if row['prof_desde_cm'] is not None and row['prof_hasta_cm'] is not None and row['prof_desde_cm']>row['prof_hasta_cm']:raise ValueError('La profundidad desde debe ser menor o igual a hasta (cm).')
    if text(row.get('fecha')):
        try:date.fromisoformat(row['fecha'])
        except ValueError:raise ValueError('Usa la fecha AAAA-MM-DD.')
    for k in CATEGORIES:row[k]=count(row.get(k),CATEGORIES[k])
    original=count(row.get('cantidad_original'));values=[v for k,v in row.items() if k in CATEGORIES and v is not None]
    described=sum(values) if values else None
    row.update(cantidad_original=original,suma_desglose=described,diferencia=None if original is None or described is None else described-original)
    if collection!='fragmentaria':
        # Las colecciones especializadas conservan sus atributos y unidades; el material no se infiere automáticamente.
        usable_quantity=original if original is not None else None
        state='Conteo declarado' if usable_quantity is not None else 'Pendiente de conteo'
    elif described is None:usable_quantity=None;state='Pendiente de desglose'
    elif original is None:usable_quantity=described;state='Cantidad derivada de desglose explícito'
    elif original==described:usable_quantity=original;state='Coincide'
    else:usable_quantity=None;state='Diferencia de cantidad'
    if text(row.get('duplicado_de')):usable_quantity=None;state='Duplicado de artefacto'
    row['cantidad_estado']=state;row['cantidad_analisis']=usable_quantity;row['usar_analisis']='Sí' if usable_quantity is not None else 'No'
    # Las cantidades descritas se conservan aunque no sean utilizables.
    row['desglose_declarado']={k:row[k] for k in CATEGORIES}
    if usable_quantity is None:
        for k in CATEGORIES:row[k]=None
    return row

def add(request):
    row=prepare_manual(request['record'])
    with database(request['database']) as con,con:
        cancelled(request)
        if text(row.get('duplicado_de')) and not con.execute('SELECT 1 FROM records WHERE id=?',(row['duplicado_de'],)).fetchone():raise ValueError('El registro original del duplicado no existe en este inventario.')
        put_record(con,row,request['record']);cancelled(request)
    return snapshot(request)

def snapshot(request):
    with database(request['database']) as con:
        rows=[json.loads(x[0]) for x in con.execute('SELECT data FROM records ORDER BY rowid')]
        metadata={k:json.loads(v) for k,v in con.execute('SELECT key,value FROM metadata')}
        containers={kind:n for kind,n in con.execute('SELECT kind,COUNT(*) FROM containers GROUP BY kind')}
    stats=statistics(rows);query=text(request.get('query')).casefold();pending=request.get('pendingOnly',False)
    shown=[r for r in rows if (not query or query in dumps(r).casefold()) and (not pending or not usable(r))]
    page=max(0,int(request.get('page',0)));page_size=100
    return {'database':request['database'],'metadata':metadata,'stats':stats,'containers':containers,'rows':shown[page*page_size:(page+1)*page_size],
            'filtered':len(shown),'page':page,'pages':max(1,math.ceil(len(shown)/page_size))}

def safe_cell(value):
    if isinstance(value,(dict,list)):return dumps(value)
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','＝','＋','－','＠')):return "'"+value
    return value

def export(request):
    output=Path(request['output']);fmt=output.suffix.lower()
    if fmt not in ('.csv','.xlsx'):raise ValueError('Exporta CSV o Excel XLSX.')
    with database(request['database']) as con:
        rows=[json.loads(x[0]) for x in con.execute('SELECT data FROM records ORDER BY rowid')]
        metadata={k:json.loads(v) for k,v in con.execute('SELECT key,value FROM metadata')}
        sources=list(con.execute('SELECT path,sha256,imported_at FROM sources'))
    if not rows:raise ValueError('No hay registros para exportar.')
    headers=list(dict.fromkeys(k for row in rows for k in row));stats=statistics(rows)
    # Reserva exclusiva de salida y archivo temporal propio: la cancelación o el error sólo elimina nuestra reserva.
    with output.open('xb'):pass
    temp=None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent,suffix=fmt,delete=False) as f:temp=Path(f.name)
        if fmt=='.csv':
            with temp.open('w',encoding='utf-8-sig',newline='') as f:
                writer=csv.writer(f);writer.writerow([safe_cell(h) for h in headers])
                for row in rows:cancelled(request);writer.writerow([safe_cell(row.get(k)) for k in headers])
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill,Alignment
            from openpyxl.utils import get_column_letter
            book=Workbook();book.remove(book.active)
            def sheet(name,head,values):
                ws=book.create_sheet(name);ws.append([safe_cell(h) for h in head])
                for row in values:cancelled(request);ws.append([safe_cell(v) for v in row])
                ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
                for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='12685F');cell.alignment=Alignment(wrap_text=True)
                ws.row_dimensions[1].height=32
                for i,h in enumerate(head,1):ws.column_dimensions[get_column_letter(i)].width=min(38,max(14,len(h)+2))
            sheet('Resumen',['Indicador','Valor'],[(k,stats[k]) for k in ('records','usable','excluded','pending','duplicates','quantity')]+[(k,v) for k,v in CATEGORIES.items()]+[('Nota','Las cantidades analíticas excluyen pendientes y duplicados. No equivale a formulario oficial MNCR.')])
            group_headers=['operacion','pozo','nivel_std','registros','cantidad',*CATEGORIES]
            sheet('Analisis',group_headers,[[r.get(k) for k in group_headers] for r in stats['groups']])
            sheet('Inventario',headers,[[r.get(k) for k in headers] for r in rows])
            sheet('Pendientes',headers,[[r.get(k) for k in headers] for r in rows if not usable(r)])
            sheet('Procedencia',['Ruta','SHA256','Fecha de importación'],sources)
            sheet('Proyecto',['Campo','Valor'],list(metadata.items()))
            sheet('MNCR - requisitos',['Colección','Formulario','Estado'],[
                ['Fragmentaria','Anexo 3: presencia X, no conteo','Usar plantilla digital oficial; no alterar columnas.'],
                ['Artefactos','Anexo 4: número, naturaleza, dimensiones, clasificación','Completar datos específicos.'],
                ['Humanos','Anexo 5: MH, individuo, sexo, edad, anatomía','No mezclar con fragmentaria.'],
                ['Fauna','Anexo 6: MF, anatomía, clase, nombres','No mezclar con fragmentaria.'],
                ['Botánica','Anexo 7: MFL, descripción, peso en gramos, nombres','No mezclar con fragmentaria.'],
                ['Carbón','Anexo 8: MC14 campo/laboratorio, peso y radiometría','No mezclar con fragmentaria.'],
                ['Fuente','Manual MNCR 2016, pp.55–58 y anexos3–8','Preparación de datos; no certificación de recepción.']])
            book.save(temp);book.close()
        cancelled(request);os.replace(temp,output);temp=None
    except Exception:
        output.unlink(missing_ok=True);raise
    finally:
        if temp:temp.unlink(missing_ok=True)
    return {'output':str(output),'sha256':digest(output),'records':len(rows),'stats':stats}

def dispatch(request):
    operations={'create':create,'open':snapshot,'list':snapshot,'preview':preview,'import':import_source,'add':add,'export':export}
    if request.get('op') not in operations:raise ValueError('Operación de inventario no compatible.')
    cancelled(request)
    try:return operations[request['op']](request)
    except sqlite3.IntegrityError as exc:raise ValueError('Importación duplicada o identificador ya utilizado. No se aplicaron cambios.') from exc

def main():
    try:
        request=json.loads(sys.stdin.buffer.read(2*1024*1024+1));result=dispatch(request)
        emit({'ok':True,'result':result})
    except Exception as exc:emit({'ok':False,'error':str(exc)})
if __name__=='__main__':main()
