"""Lectura y escritura local acotada de XLSX; los cálculos geodésicos permanecen en el motor C++ PROJ."""
from pathlib import Path
import csv
import io
import json
import math
import os
import sys
import tempfile
import zipfile


def cell_text(value):
    if value is None: return ""
    if isinstance(value, float) and not math.isfinite(value): raise ValueError("file.xlsx_invalid")
    return str(value)


def read_sheet(request):
    from openpyxl import load_workbook
    path = Path(request["path"])
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("file.xlsx_limit")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
            raise ValueError("file.xlsx_limit")
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        names = values.sheetnames
        if not request.get("sheet"):
            return {"sheets": names, "path": str(path)}
        sheet = request["sheet"]
        if sheet not in names: raise ValueError("file.xlsx_sheet")
        ws = values[sheet]
        if (ws.max_column or 0) > 128 or (ws.max_row or 0) > 10001:
            raise ValueError("file.xlsx_limit")
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, delimiter="\t", lineterminator="\n")
        cached_formulas = 0
        for count, (row, formula_row) in enumerate(zip(ws.iter_rows(values_only=True), formulas[sheet].iter_rows()), 1):
            if count > 10001 or len(row) > 128: raise ValueError("file.xlsx_limit")
            for value, cell in zip(row, formula_row):
                if cell.data_type == "f":
                    if value is None: raise ValueError("file.xlsx_formula_cache")
                    cached_formulas += 1
                if cell.data_type == "e": raise ValueError("file.xlsx_cell_error")
            writer.writerow([cell_text(v) for v in row])
            if buffer.tell() > 8 * 1024 * 1024: raise ValueError("input.too_large")
        return {"sheets": names, "sheet": sheet, "path": str(path), "text": buffer.getvalue(), "cachedFormulas": cached_formulas}
    finally:
        values.close()
        if formulas: formulas.close()


def text_cell(ws, value):
    from openpyxl.cell import WriteOnlyCell
    text = cell_text(value)
    if len(text) > 32767: raise ValueError("file.xlsx_limit")
    cell = WriteOnlyCell(ws, value=text)
    cell.data_type = "s"  # La entrada que empieza por =, +, - o @ siempre se trata como texto literal.
    return cell


def write_sheet(request):
    from openpyxl import Workbook
    destination = Path(request["path"])
    if destination.suffix.lower() != ".xlsx": raise ValueError("file.xlsx_extension")
    if os.path.lexists(destination): raise FileExistsError("file.exists")
    rows = list(csv.reader(io.StringIO(request["csv"])))
    if len(rows) > 10001 or not rows: raise ValueError("file.xlsx_limit")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Coordenadas")
    ws.freeze_panes = "A2"
    technical_start = len(rows[0]) - 20
    numeric = {technical_start + 3, technical_start + 4, technical_start + 15, technical_start + 17}
    for index, row in enumerate(rows):
        values = []
        for column, value in enumerate(row):
            if index and column in numeric and value:
                number = float(value)
                if not math.isfinite(number): raise ValueError("file.xlsx_invalid")
                values.append(number)
            else: values.append(text_cell(ws, value))
        ws.append(values)
    provenance = wb.create_sheet("Procedencia")
    provenance.append([text_cell(provenance, "Información"), text_cell(provenance, "Valor")])
    provenance.append([text_cell(provenance, "Precisión"), text_cell(provenance, "Excel presenta números con su propia precisión. El JSON de procedencia conserva la representación completa del motor.")])
    metadata = dict(request["result"])
    metadata.pop("rows", None)
    metadata["inputSpreadsheet"] = request.get("inputSpreadsheet", {})
    encoded = json.dumps(metadata, ensure_ascii=False, allow_nan=False)
    for i in range(0, len(encoded), 30000):
        provenance.append([text_cell(provenance, f"JSON {i // 30000 + 1}"), text_cell(provenance, encoded[i:i+30000])])
    raw = wb.create_sheet("Resultados_JSON")
    for row in request["result"].get("rows", []):
        encoded = json.dumps(row, ensure_ascii=False, allow_nan=False)
        raw.append([text_cell(raw, encoded[i:i+30000]) for i in range(0, len(encoded), 30000)])
    fd, temporary = tempfile.mkstemp(prefix=".archaeohub-xlsx-", suffix=".tmp", dir=destination.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        with open(temporary, "rb+") as stream: os.fsync(stream.fileno())
        os.link(temporary, destination)  # Publicación atómica que nunca sustituye un nombre existente.
    finally:
        Path(temporary).unlink(missing_ok=True)
        wb.close()
    return {"path": str(destination), "rows": len(rows) - 1}


def dispatch(request):
    if request.get("schemaVersion") != 1: raise ValueError("engine.protocol")
    op = request.get("op")
    if op == "readSpreadsheet": return read_sheet(request)
    if op == "writeSpreadsheet": return write_sheet(request)
    raise ValueError("engine.operation")


def main():
    try:
        raw = sys.stdin.buffer.read(32 * 1024 * 1024 + 1)
        if len(raw) > 32 * 1024 * 1024: raise ValueError("input.too_large")
        result = dispatch(json.loads(raw))
        reply = {"type": "result", "ok": True, **result}
        code = 0
    except Exception as exc:
        message = str(exc)
        reply = {"type": "error", "ok": False, "code": "file.exists" if isinstance(exc, FileExistsError) else message if message.startswith(("file.", "input.", "engine.")) else "file.xlsx_invalid"}
        code = 1
    sys.stdout.buffer.write((json.dumps(reply, ensure_ascii=False, allow_nan=False) + "\n").encode("utf-8"))
    sys.stdout.buffer.flush()
    return code


if __name__ == "__main__": sys.exit(main())
