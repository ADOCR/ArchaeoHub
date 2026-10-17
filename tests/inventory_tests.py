from pathlib import Path
import unittest,tempfile,sys,json,sqlite3
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'modules/inventory'))
import inventory_worker as w
class InventoryTests(unittest.TestCase):
 def test_untrusted_headers_remain_literal_in_exports(self):
  import csv
  from openpyxl import load_workbook
  self.request('add',record={'operacion':'2','q_cer':3,'=1+1':'literal','comentarios':'＠SUM(1)'})
  for extension in ('csv','xlsx'):
   path=self.root/('safe.'+extension);self.request('export',output=str(path))
   if extension=='csv':
    with path.open(encoding='utf-8-sig',newline='') as f:values=list(csv.reader(f))
    self.assertIn("'=1+1",values[0]);self.assertNotIn('=1+1',values[0])
   else:
    book=load_workbook(path,data_only=False)
    self.assertFalse(any(c.data_type=='f' for sheet in book for row in sheet for c in row));book.close()
  self.assertIn('=1+1',self.request('open')['rows'][0])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.db=self.root/'Inventario.ahinv'
  w.create({'database':str(self.db),'metadata':{'project':'Sintético','site':'QA'}})
 def tearDown(self):self.tmp.cleanup()
 def request(self,op,**kw):return w.dispatch({'database':str(self.db),'op':op,**kw})
 def source(self,rows):
  p=self.root/'source.json';p.write_text(json.dumps({'inventario':rows}),encoding='utf-8');return str(p)
 def test_mismatch_null_not_zero(self):
  result=self.request('add',record={'operacion':'2','cantidad_original':10,'q_cer':8})
  row=result['rows'][0];self.assertEqual(row['cantidad_original'],10);self.assertEqual(row['diferencia'],-2);self.assertIsNone(row['cantidad_analisis']);self.assertIsNone(row['q_cer']);self.assertEqual(row['desglose_declarado']['q_cer'],8);self.assertEqual(result['stats']['quantity'],0)
 def test_derive_only_explicit_breakdown(self):
  result=self.request('add',record={'operacion':'2','q_cer':0,'q_lit':4,'nivel_std':'2-3'})
  row=result['rows'][0];self.assertEqual(row['cantidad_analisis'],4);self.assertIsNone(row['cantidad_original']);self.assertIsNone(row['prof_desde_cm']);self.assertEqual(row['nivel_std'],'2-3')
 def test_original_without_breakdown_remains_pending(self):
  self.assertIsNone(self.request('add',record={'operacion':'2','cantidad_original':4})['rows'][0]['cantidad_analisis'])
 def test_container_counts_are_distinct(self):
  for _ in range(2):result=self.request('add',record={'operacion':'2','q_cer':3,'caja':'01','bolsa_grande':'A','bolsita':'1'})
  self.assertEqual(result['containers'],{'box':1,'large_bag':1,'small_bag':1});self.assertEqual(result['stats']['quantity'],6)
 def test_preview_no_write_duplicate_rollback(self):
  rows=[{'registro':'R1','cantidad_original':7,'cantidad_analisis':7,'usar_analisis':'Sí','q_cer':7}]
  src=self.source(rows);p=self.request('preview',source=src);self.assertEqual(self.request('open')['stats']['records'],0)
  self.request('import',source=src,sourceDigest=p['sourceDigest'])
  with self.assertRaises(ValueError):self.request('import',source=src,sourceDigest=p['sourceDigest'])
  self.assertEqual(self.request('open')['stats']['quantity'],7)
 def test_import_collision_rolls_back_all_rows(self):
  self.request('add',record={'registro':'R1','operacion':'2','q_cer':7})
  src=self.source([{'registro':rid,'cantidad_original':2,'cantidad_analisis':2,'usar_analisis':'Sí'} for rid in ['R2','R1']]);p=self.request('preview',source=src)
  self.assertEqual(p['collisions'],['R1'])
  with self.assertRaises(ValueError):self.request('import',source=src,sourceDigest=p['sourceDigest'])
  self.assertEqual(self.request('open')['stats']['records'],1)
 def test_changed_source_refused(self):
  src=self.source([{'registro':'R','cantidad_original':1,'cantidad_analisis':1,'usar_analisis':'Sí'}]);p=self.request('preview',source=src);Path(src).write_text('[]')
  with self.assertRaises(ValueError):self.request('import',source=src,sourceDigest=p['sourceDigest'])
  self.assertEqual(self.request('open')['stats']['records'],0)
 def test_cancel_rollback_and_export_cleanup(self):
  self.request('add',record={'operacion':'2','q_cer':2});flag=self.root/'cancel';flag.touch();output=self.root/'cancelled.csv'
  with self.assertRaises(ValueError):self.request('export',output=str(output),cancelFile=str(flag))
  self.assertFalse(output.exists())
 def test_no_clobber(self):
  before=w.digest(self.db)
  with self.assertRaises(FileExistsError):w.create({'database':str(self.db),'metadata':{'site':'X','project':'Y'}})
  self.assertEqual(before,w.digest(self.db));self.request('add',record={'operacion':'2','q_cer':2});out=self.root/'existing.csv';out.write_text('USER')
  with self.assertRaises(FileExistsError):self.request('export',output=str(out))
  self.assertEqual(out.read_text(),'USER')
 def test_excel_roundtrip_safe_text(self):
  self.request('add',record={'operacion':'2','q_cer':2,'comentarios_laboratorio':'=1+1','pozo':'001'})
  out=self.root/'Analysis.xlsx';self.request('export',output=str(out));from openpyxl import load_workbook
  book=load_workbook(out,data_only=False);sheet=book['Inventario'];rows=list(sheet.values);data=dict(zip(rows[0],rows[1]));self.assertEqual(data['pozo'],'001');self.assertEqual(data['comentarios_laboratorio'],"'=1+1");book.close()
  source_hash=w.digest(out);p=self.request('preview',source=str(out));self.assertEqual(p['quantity'],2);self.assertEqual(w.digest(out),source_hash)
 def test_invalid_depth_count_date(self):
  for fields in ({'q_cer':1.5},{'prof_desde_cm':40,'prof_hasta_cm':20},{'fecha':'ayer'},{'q_cer':-1}):
   with self.assertRaises(ValueError):self.request('add',record={'operacion':'2',**fields})
  self.assertEqual(self.request('open')['stats']['records'],0)
 def test_duplicate_excluded_and_preserved(self):
  src=self.source([{'registro':'A','cantidad_original':3,'cantidad_analisis':3,'usar_analisis':'Sí'}, {'registro':'B','cantidad_original':3,'cantidad_analisis':3,'usar_analisis':'Sí','duplicado_de':'A'}]);p=self.request('preview',source=src);self.assertEqual(p['quantity'],3);self.assertEqual(p['duplicates'],1)
 def test_unrecognized_database_read_only(self):
  other=self.root/'other.sqlite';con=sqlite3.connect(other);con.execute('create table important(data)');con.close();before=w.digest(other)
  with self.assertRaises(ValueError):w.snapshot({'database':str(other)})
  self.assertEqual(before,w.digest(other))
 def test_excel_expansion_limit_before_parser(self):
  import zipfile
  p=self.root/'Oversized.xlsx'
  with zipfile.ZipFile(p,'w') as archive:
   for i in range(4097):archive.writestr(str(i),'')
  with self.assertRaisesRegex(ValueError,'expandido'):self.request('preview',source=str(p))
if __name__=='__main__':unittest.main()
