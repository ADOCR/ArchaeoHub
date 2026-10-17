"""Sólo datos sintéticos. Ejecuta con el Python de desarrollo del catálogo (Pillow/openpyxl)."""
import ast
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "modules/catalog/python"))
from PIL import Image
from openpyxl import load_workbook
from catalog_worker import Session
from catalog_core import Catalog, CatalogLock, FIELDS, LEGACY_FIELDS_V3
from catalog_project import inspect_project


class CatalogBridgeTests(unittest.TestCase):
    def test_catalog_forged_dimensions_are_rejected_before_iteration(self):
        import zipfile
        from catalog_workbook import load_workbook as bounded_load
        file = self.root / "oversized.xlsx"
        book = self.session.catalog.wb
        book.save(file)
        with zipfile.ZipFile(file) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'<dimension ref="A1:U1"', b'<dimension ref="A1:XFD1048576"')
        # El ancho del encabezado varía entre versiones anteriores; sustituye la dimensión de forma genérica.
        import re
        entries["xl/worksheets/sheet1.xml"] = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', entries["xl/worksheets/sheet1.xml"])
        with zipfile.ZipFile(file, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items(): archive.writestr(name, data)
        with self.assertRaisesRegex(ValueError, "100.000"):
            bounded_load(file, read_only=True)
        self.assert_originals()
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ArchaeoHub-catalog-")
        self.root = Path(self.temp.name)
        self.source, self.output, self.cache = [self.root / p for p in ("Originales ñ", "Trabajo", "Cache")]
        self.source.mkdir(); self.output.mkdir()
        for name in ("foto A.jpg", "foto B.jpg"):
            image = Image.new("RGB", (64, 48), (23, 100, 82))
            exif = Image.Exif(); exif[274] = 6; exif[36867] = "2026:01:02 03:04:05"
            image.save(self.source / name, exif=exif)
        self.hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.source.iterdir()}
        self.session = Session()
        self.open_request = {"schemaVersion": 1, "op": "open", "source": str(self.source), "output": str(self.output), "cache": str(self.cache), "site": "TEST-ONLY", "operation": "2"}
        self.form = {"operation": "2", "unitType": "Pozo", "unit": "7", "level": "1", "subunitType": "Cuadro", "subunit": "A", "view": "Perfil norte", "notes": "Descripción ñ"}
        self.result = self.session.dispatch(self.open_request)
        self.id = self.result["items"][0]["id"]

    def tearDown(self):
        self.session.close(); self.temp.cleanup()

    def request(self, op, **kwargs):
        return self.session.dispatch({"schemaVersion": 1, "op": op, "id": self.id, **kwargs})

    def assert_originals(self):
        self.assertEqual(self.hashes, {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.source.iterdir()})

    def test_preview_orientation_and_original_hash(self):
        selection = self.request("select")["selection"]
        from urllib.parse import urlparse, unquote
        path = Path(unquote(urlparse(selection["preview"]["url"]).path).lstrip("/") if os.name == "nt" else unquote(urlparse(selection["preview"]["url"]).path))
        with Image.open(path) as image:
            self.assertEqual(image.size, (48, 64))
        self.assert_originals()

    def test_save_correct_withdraw_and_reopen(self):
        predicted = self.request("propose", form=self.form)["proposedName"]
        self.request("save", form=self.form)
        copy = self.output / "Entrega_Museo/Fotos" / predicted
        original = self.session.paths[self.id]
        self.assertEqual(copy.read_bytes(), original.read_bytes())
        wb = load_workbook(self.session.catalog.xlsx_path)
        try:
            self.assertEqual([c.value for c in wb.active[1]], FIELDS)
            self.assertEqual(wb.active.cell(2, FIELDS.index("Observaciones") + 1).value, "Descripción ñ")
        finally:
            wb.close()
        self.form["unit"] = "9"
        changed = self.request("correct", form=self.form)
        self.assertEqual(changed["selection"]["id"], self.id)
        self.assertFalse(copy.exists())
        self.request("withdraw", confirmed=True)
        self.assertEqual(len(self.session.catalog.rows), 0)
        self.assertEqual(self.session.catalog.status_of(original), "skipped")
        self.assert_originals()
        self.session.close(); self.session.dispatch(self.open_request)
        self.assertEqual(len(self.session.catalog.rows), 0)

    def test_skip_recover_filters_preserve_state(self):
        self.request("skip", skipped=True)
        page = self.request("page", filter="Omitidas")
        self.assertEqual(page["matched"], 1)
        self.request("skip", skipped=False)
        self.assertEqual(self.request("page", filter="Omitidas")["matched"], 0)
        self.assert_originals()

    def test_collision_does_not_overwrite(self):
        name = self.request("propose", form=self.form)["proposedName"]
        target = self.session.catalog.photos / name
        target.write_bytes(b"KEEP")
        with self.assertRaises(FileExistsError): self.request("save", form=self.form)
        self.assertEqual(target.read_bytes(), b"KEEP")
        self.assertEqual(len(self.session.catalog.rows), 0)
        self.assert_originals()

    def test_discard_requires_confirmation_and_is_recoverable(self):
        original = self.session.paths[self.id]
        with self.assertRaises(ValueError): self.request("discard")
        self.assertTrue(original.exists())
        self.request("discard", confirmed=True)
        self.assertFalse(original.exists())
        backups = list((self.output / "Descartadas_catalogador").glob("*/*.jpg"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(hashlib.sha256(backups[0].read_bytes()).hexdigest(), self.hashes[original.name])
        metadata = json.loads((backups[0].parent / "origen.json").read_text(encoding="utf-8"))
        self.assertEqual(Path(metadata["original"]), original)

    def test_lock_failure_keeps_current_session(self):
        other = Session()
        with self.assertRaises(ValueError): other.dispatch(self.open_request)
        self.assertEqual(self.request("page")["matched"], 2)
        other.close()

    def test_wrong_site_keeps_existing_rows(self):
        self.request("save", form=self.form)
        before = self.session.catalog.xlsx_path.read_bytes()
        with self.assertRaises(ValueError): self.session.dispatch(self.open_request | {"site": "OTHER"})
        self.assertEqual(self.session.catalog.xlsx_path.read_bytes(), before)

    def test_unknown_photo_cannot_write(self):
        with self.assertRaises(ValueError): self.request("save", id="../../file", form=self.form)

    def test_report_is_separate_and_refuses_overwrite(self):
        self.request("save", form=self.form)
        before = self.session.catalog.xlsx_path.read_bytes()
        destination = self.output / "GPS.xlsx"
        result = self.request("exportGps", destination=str(destination))
        self.assertEqual(result["report"]["photos"], 1)
        self.assertEqual(before, self.session.catalog.xlsx_path.read_bytes())
        with self.assertRaises(FileExistsError): self.request("exportGps", destination=str(destination))
        with self.assertRaises(ValueError): self.request("exportGps", destination=str(self.session.catalog.delivery / "GPS.xlsx"))
        self.assert_originals()

    def test_manual_gps_modifies_only_copy_with_backup(self):
        self.request("save", form=self.form)
        reply=self.request("manualGps", manualGps=["10.1", "-84.2", "450", "5"])
        self.assertAlmostEqual(reply["selection"]["copyGps"]["latitude"],10.1,places=6)
        self.assertEqual(reply["selection"]["gps"],{})
        self.assertEqual(len(list((self.output / "Respaldos_GPS_manual").iterdir())), 1)
        self.assert_originals()
        with self.assertRaises(ValueError): self.request("manualGps", manualGps=["10.2", "-84.2", "", ""])

    def test_bad_open_keeps_existing_session(self):
        with self.assertRaises(ValueError): self.session.dispatch(self.open_request | {"output": str(self.source)})
        self.assertEqual(self.request("page")["matched"], 2)

    def test_legacy_schema_needs_confirmation(self):
        self.session.close()
        xlsx = self.output / "Entrega_Museo/Catalogo_fotografico_pozos.xlsx"
        from openpyxl import Workbook
        wb = Workbook(); wb.active.append(LEGACY_FIELDS_V3); wb.active.append(["TEST-ONLY","7","1",1,"old.jpg","old.jpg","","","",""]); wb.save(xlsx);wb.close()
        before = xlsx.read_bytes()
        reply = self.session.dispatch(self.open_request)
        self.assertTrue(reply["migrationRequired"])
        self.assertEqual(before, xlsx.read_bytes())
        # Rechazar la acción debe liberar el bloqueo.
        lock = CatalogLock(self.output);lock.close()

    def test_committed_warning_requires_reopening(self):
        from unittest.mock import patch
        from catalog_core import SaveCommittedWarning
        with patch.object(self.session.catalog, "save", side_effect=SaveCommittedWarning("La copia quedó confirmada; reabrir.")):
            result = self.request("save", form=self.form)
        self.assertTrue(result["state"]["recoveryRequired"])
        self.assertTrue(result["warning"])
        with self.assertRaises(ValueError): self.request("skip", skipped=True)
        self.session.dispatch(self.open_request)
        self.assertFalse(self.session.state()["recoveryRequired"])

    def test_load_existing_excel_without_site_or_operation(self):
        self.request("save", form=self.form)
        self.session.close()
        files = {str(p.relative_to(self.output)): p.read_bytes() for p in self.output.rglob("*") if p.is_file()}
        reply = self.session.dispatch({"schemaVersion": 1, "op": "load", "location": str(self.output / "Entrega_Museo/Catalogo_fotografico_pozos.xlsx"), "source": str(self.source), "cache": str(self.cache)})
        self.assertEqual(reply["state"]["registered"], 1)
        self.assertEqual(reply["project"]["site"], "TEST-ONLY")
        self.assertEqual(reply["project"]["operation"], "2")
        self.assertFalse(reply["project"]["needsSite"])
        self.session.close()  # El bloqueo activo impide deliberadamente las lecturas en Windows.
        self.assertEqual(files, {str(p.relative_to(self.output)): p.read_bytes() for p in self.output.rglob("*") if p.is_file()})
        self.assert_originals()

    def test_load_requires_only_unknown_source_and_does_not_create(self):
        self.request("save", form=self.form)
        self.session.close()
        before = {str(p): p.read_bytes() for p in self.output.rglob("*") if p.is_file()}
        reply = self.session.dispatch({"schemaVersion": 1, "op": "load", "location": str(self.output), "cache": str(self.cache)})
        self.assertTrue(reply["needsInput"])
        self.assertTrue(reply["project"]["needsSource"])
        self.assertEqual(reply["project"]["suggestedSource"], str(self.source))
        self.assertEqual(before, {str(p): p.read_bytes() for p in self.output.rglob("*") if p.is_file()})
        empty = self.root / "No es catálogo"; empty.mkdir()
        with self.assertRaises(ValueError): self.session.dispatch({"schemaVersion": 1, "op": "load", "location": str(empty), "cache": str(self.cache)})
        self.assertEqual(list(empty.iterdir()), [])

    def test_load_multiple_operations_does_not_invent_default(self):
        self.request("save", form=self.form)
        second = next(i for i in self.session.paths if i != self.id)
        self.request("save", id=second, form=self.form | {"operation": "3"})
        self.session.close()
        reply = self.session.dispatch({"schemaVersion": 1, "op": "load", "location": str(self.output / "estado_catalogador.json"), "source": str(self.source), "cache": str(self.cache)})
        self.assertEqual(reply["state"]["registered"], 2)
        self.assertEqual(reply["project"]["operations"], ["2", "3"])
        self.assertEqual(reply["project"]["operation"], "")
        self.assert_originals()

    def test_save_next_manual_gps_and_resume_pending(self):
        saved = self.session.paths[self.id]
        reply = self.request("save", form=self.form, advance=True, manualGps=["10.1", "-84.2", "450", "5"])
        self.assertNotEqual(reply["selection"]["id"], self.id)
        next_id = reply["selection"]["id"]
        self.assertFalse(reply["selection"]["registered"])
        copy = self.session.catalog.photos / self.session.catalog.rows[0]["Nombre_archivo"]
        from georreferenciacion import read_gps
        point, _ = read_gps(copy)
        self.assertAlmostEqual(point.latitude, 10.1, places=6)
        self.assert_originals()
        self.session.close()
        reopened = self.session.dispatch(self.open_request)
        self.assertEqual(reopened["selection"]["id"], next_id)
        self.assertEqual(self.session.catalog.status_of(saved), "registered")

    def test_registered_discard_is_recoverable(self):
        self.request("save", form=self.form)
        self.request("discard", confirmed=True)
        self.assertEqual(len(self.session.catalog.rows), 0)
        self.assertEqual(len(list((self.output / "Descartadas_catalogador").glob("*/*.jpg"))), 1)

    def test_create_refuses_existing_catalog(self):
        self.request("save", form=self.form)
        before = self.session.catalog.xlsx_path.read_bytes()
        with self.assertRaises(ValueError): self.session.dispatch(self.open_request | {"create": True})
        self.assertEqual(before, self.session.catalog.xlsx_path.read_bytes())

    def test_next_gps_finds_registered_review_photo(self):
        self.request("save", form=self.form)
        reply = self.request("nextGps")
        self.assertEqual(reply["selection"]["id"], self.id)
        self.assertEqual(reply["filter"], "Todas")

    def test_navigation_crosses_thumbnail_page_boundary(self):
        self.session.close()
        for i in range(65):
            (self.source / f"extra-{i:03}.jpg").write_bytes((self.source / "foto A.jpg").read_bytes())
        reply = self.session.dispatch(self.open_request)
        current = reply["items"][59]["id"]
        next_page = self.request("navigate", id=current, delta=1, filter="Todas", query="")
        self.assertEqual(next_page["offset"], 60)
        self.assertEqual(next_page["selection"]["id"], next_page["items"][0]["id"])
        previous = self.request("navigate", id=next_page["selection"]["id"], delta=-1)
        self.assertEqual(previous["selection"]["id"], current)

    def test_failed_save_does_not_advance_or_change_originals(self):
        from unittest.mock import patch
        with patch.object(self.session.catalog, "_write_workbook", side_effect=PermissionError("Excel bloqueado")):
            with self.assertRaises(PermissionError): self.request("save", form=self.form, advance=True)
        self.assertEqual(self.session.catalog.status_of(self.session.paths[self.id]), "pending")
        self.assertEqual(len(self.session.catalog.rows), 0)
        self.assert_originals()

    def test_partial_registered_discard_reports_committed_withdrawal(self):
        from unittest.mock import patch
        self.request("save", form=self.form)
        original = self.session.paths[self.id]
        with patch.object(self.session.catalog, "discard_unregistered", side_effect=PermissionError("Unidad ocupada")):
            reply = self.request("discard", confirmed=True)
        self.assertTrue(reply["committed"])
        self.assertTrue(reply["state"]["recoveryRequired"])
        self.assertEqual(reply["state"]["registered"], 0)
        self.assertTrue(reply["selection"]["skipped"])
        self.assertTrue(original.is_file())
        self.assertTrue(reply["warning"])
        self.assert_originals()


if __name__ == "__main__":
    unittest.main()
