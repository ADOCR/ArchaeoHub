"""Las pruebas XLSX validan valores, hojas, errores y conservación sin datos de campo."""
from pathlib import Path
import csv
import io
import json
import sys
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"modules/coordinates"))
from spreadsheet_worker import dispatch

class SpreadsheetContracts(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.path=self.root/"Coordenadas ñ.xlsx"
        wb=Workbook();ws=wb.active;ws.title="Campo";ws.append(["id","x","y","nota"]);ws.append(["001",-84.0,10.0,"= literal tratado como texto"])
        ws["D2"].data_type="s";ws=wb.create_sheet("Otras");ws.append(["id","x","y"]);ws.append(["002","-84,1","10,1"]);wb.save(self.path);wb.close()
        self.before=self.path.read_bytes()
    def tearDown(self):self.temp.cleanup()
    def rpc(self,op,**kw):return dispatch({"schemaVersion":1,"op":op,**kw})
    def test_sheets_and_literal_values_preserved(self):
        result=self.rpc("readSpreadsheet",path=str(self.path))
        self.assertEqual(result["sheets"],["Campo","Otras"])
        result=self.rpc("readSpreadsheet",path=str(self.path),sheet="Campo")
        rows=list(csv.reader(io.StringIO(result["text"]),delimiter="\t"))
        self.assertEqual(rows[1],["001","-84","10","= literal tratado como texto"])
        self.assertEqual(self.before,self.path.read_bytes())
    def test_formulas_without_cached_values_fail_without_changes(self):
        path=self.root/"Formula.xlsx";wb=Workbook();wb.active.append(["=1+2",4]);wb.save(path);wb.close();before=path.read_bytes()
        with self.assertRaisesRegex(ValueError,"file.xlsx_formula_cache"):self.rpc("readSpreadsheet",path=str(path),sheet="Sheet")
        self.assertEqual(before,path.read_bytes())
    def test_export_numeric_results_literal_ids_and_provenance_no_overwrite(self):
        out=self.root/"Entrega.xlsx";headers=["original"]+["c"+str(i) for i in range(20)]
        values=["=HYPERLINK(\"malicious\")"]+[""]*20
        values[4]="500000.125";values[5]="1105412.375";values[16]="0";values[18]="0.1"
        buffer=io.StringIO();writer=csv.writer(buffer);writer.writerows([headers,values])
        req={"path":str(out),"csv":buffer.getvalue(),"result":{"rows":[{"id":"001","resultX":500000.125}],"source":{"code":"EPSG:5365"},"target":{"code":"EPSG:5367"},"operations":[{"name":"Verificado"}]}}
        self.rpc("writeSpreadsheet",**req)
        wb=load_workbook(out,data_only=False)
        try:
            self.assertEqual(wb.sheetnames,["Coordenadas","Procedencia","Resultados_JSON"])
            ws=wb["Coordenadas"];self.assertEqual(ws.cell(2,1).data_type,"s");self.assertEqual(ws.cell(2,5).value,500000.125);self.assertEqual(ws.cell(2,5).data_type,"n")
            self.assertEqual(json.loads(wb["Resultados_JSON"].cell(1,1).value)["id"],"001")
        finally:wb.close()
        before=out.read_bytes()
        with self.assertRaises(FileExistsError):self.rpc("writeSpreadsheet",**req)
        self.assertEqual(before,out.read_bytes());self.assertEqual(self.before,self.path.read_bytes())
    def test_oversize_and_unknown_sheet_fail(self):
        with self.assertRaisesRegex(ValueError,"file.xlsx_sheet"):self.rpc("readSpreadsheet",path=str(self.path),sheet="No existe")
        path=self.root/"Grande.xlsx";wb=Workbook();wb.active.cell(10002,1,"x");wb.save(path);wb.close()
        with self.assertRaisesRegex(ValueError,"file.xlsx_limit"):self.rpc("readSpreadsheet",path=str(path),sheet="Sheet")

if __name__=="__main__":unittest.main()
